import json
import io
import zipfile
from datetime import date, time
from html.parser import HTMLParser

from django.contrib.auth import authenticate, login, logout
from django.contrib.auth.models import Group
from django.contrib.auth.decorators import login_required
from django.core.exceptions import ValidationError
from django.core.validators import validate_email
from django.db import IntegrityError, transaction
from django.db.models.deletion import ProtectedError
from django.db.models import Prefetch, Q
from django.http import FileResponse, HttpResponse, HttpResponseBadRequest, JsonResponse
from django.shortcuts import get_object_or_404
from django.middleware.csrf import get_token
from django.views.decorators.csrf import ensure_csrf_cookie
from django.views.decorators.cache import never_cache
from django.views.decorators.http import require_GET, require_POST, require_http_methods
from django.views.csrf import csrf_failure as default_csrf_failure
from django.utils import timezone

from .filters import alerts_for_user, apply_section_filters, documents_for_user, sections_for_user
from .models import APISection, Alert, AlertRead, AuditLog, CalendarEvent, Campus, CampusCareer, Career, Course, CTPValidationReview, CycleTemplate, Document, DocumentVersion, ProcessStage, ProductivePartner, Semester, User, UserCampus, Visit
from .permissions import (
    ADMIN_ROLES,
    can_access_section,
    can_advance_stage,
    can_define_requirements,
    can_download_document,
    can_manage_stage,
    can_manage_section,
    can_manage_visits,
    can_propose_partner,
    can_review_document,
    can_select_partner,
    can_upload_document,
    has_role,
)
from .serializers import alert_to_dict, document_to_dict, section_to_dict, stage_to_dict, visit_to_dict
from .services import (
    advance_stage,
    add_document_version,
    delete_document_version,
    generate_due_alerts,
    log_action,
    propose_partner_for_section,
    review_document,
    resolve_stage_alerts,
    select_partner_for_section,
    update_section_requirements,
)


# Lee el cuerpo JSON de una peticion y devuelve None si no es valido.
def _json_body(request):
    if not request.body:
        return {}
    try:
        return json.loads(request.body.decode("utf-8"))
    except json.JSONDecodeError:
        return None


# Respuesta estándar para accesos no autorizados.
def _forbidden():
    return JsonResponse({"detail": "No tienes permisos para acceder a este recurso."}, status=403)


# Evita que la API devuelva la pagina HTML completa de Django cuando falla CSRF.
def csrf_failure(request, reason=""):
    accepts_json = "application/json" in request.headers.get("Accept", "")
    if request.path.startswith("/api/") or accepts_json:
        return JsonResponse(
            {
                "detail": "La sesión de seguridad vencio o cambio. Intenta nuevamente; si sigue ocurriendo, recarga la pagina e inicia sesión otra vez.",
                "code": "csrf_failed",
            },
            status=403,
        )
    return default_csrf_failure(request, reason=reason)

# Respuesta estándar para errores de validación de negocio.
def _validation_error(error):
    messages = error.messages if hasattr(error, "messages") else [str(error)]
    return JsonResponse({"detail": messages}, status=400)

DEPENDENCY_LABELS = {
    "APISection": "secciones API",
    "ProcessStage": "etapas",
    "Document": "documentos",
    "DocumentVersion": "versiones de documentos",
    "Visit": "visitas",
    "Alert": "alertas",
    "AlertRead": "lecturas de alertas",
    "AuditLog": "auditoría",
    "CTPValidationReview": "revisiones CTP",
    "CampusCareer": "carreras por sede",
    "UserCampus": "sedes asignadas a usuarios",
}


def _protected_dependency_message(error):
    protected_objects = list(getattr(error, "protected_objects", []) or [])
    counts = {}
    for item in protected_objects:
        model_name = item.__class__.__name__
        label = DEPENDENCY_LABELS.get(model_name, item._meta.verbose_name_plural or item._meta.verbose_name)
        counts[label] = counts.get(label, 0) + 1
    if not counts:
        return "El registro está en uso por otros datos del sistema. Desactívalo en lugar de eliminarlo."
    details = ", ".join(f"{count} {label}" for label, count in sorted(counts.items()))
    return f"El registro está en uso por: {details}. Desactívalo en lugar de eliminarlo."


# Obtiene la IP del cliente para registros de auditoría.
def _client_ip(request):
    forwarded_for = request.META.get("HTTP_X_FORWARDED_FOR")
    if forwarded_for:
        return forwarded_for.split(",")[0].strip()
    return request.META.get("REMOTE_ADDR")


# Convierte una fecha ISO enviada por React en objeto date de Python.
def _parse_date(value):
    if not value:
        raise ValidationError("La fecha es obligatoria.")
    try:
        return date.fromisoformat(value)
    except ValueError as exc:
        raise ValidationError("La fecha debe tener formato YYYY-MM-DD.") from exc

def _parse_optional_time(value):
    if not value:
        return None
    try:
        return time.fromisoformat(value)
    except ValueError as exc:
        raise ValidationError("La hora debe tener formato HH:MM.") from exc

class _ExcelHtmlTableParser(HTMLParser):
    def __init__(self):
        super().__init__()
        self.rows = []
        self.cell_types = []
        self.caption = ""
        self._current_row = None
        self._current_row_types = None
        self._current_cell = None
        self._current_cell_type = None
        self._in_cell = False
        self._in_caption = False
        self._caption_parts = []

    def handle_starttag(self, tag, attrs):
        if tag == "tr":
            self._current_row = []
            self._current_row_types = []
        if tag in {"td", "th"} and self._current_row is not None:
            self._current_cell = []
            self._current_cell_type = tag
            self._in_cell = True
        if tag == "caption":
            self._in_caption = True
            self._caption_parts = []

    def handle_data(self, data):
        if self._in_cell and self._current_cell is not None:
            self._current_cell.append(data)
        elif self._in_caption:
            self._caption_parts.append(data)

    def handle_endtag(self, tag):
        if tag in {"td", "th"} and self._in_cell and self._current_row is not None:
            self._current_row.append(" ".join("".join(self._current_cell).split()))
            self._current_row_types.append(self._current_cell_type)
            self._current_cell = None
            self._current_cell_type = None
            self._in_cell = False
        if tag == "tr" and self._current_row is not None:
            if any(cell for cell in self._current_row):
                self.rows.append(self._current_row)
                self.cell_types.append(self._current_row_types)
            self._current_row = None
            self._current_row_types = None
        if tag == "caption":
            self.caption = " ".join("".join(self._caption_parts).split())
            self._in_caption = False
            self._caption_parts = []


def _preview_xlsx(file_obj, max_rows=25, max_columns=12):
    from openpyxl import load_workbook

    workbook = load_workbook(file_obj, read_only=True, data_only=True)
    sheet = workbook.active
    rows = []
    for row in sheet.iter_rows(max_row=max_rows, max_col=max_columns, values_only=True):
        values = ["" if value is None else str(value) for value in row]
        if any(values):
            rows.append(values)
    workbook.close()
    return {"sheet": sheet.title, "rows": rows}


def _preview_html_excel(file_obj, max_rows=25, max_columns=12):
    content = file_obj.read()
    text = content.decode("utf-8-sig", errors="ignore")
    parser = _ExcelHtmlTableParser()
    parser.feed(text)
    rows = [row[:max_columns] for row in parser.rows[:max_rows]]
    return {"sheet": "Plantilla", "rows": rows}


def _excel_preview(version):
    file_name = version.file.name.lower()
    with version.file.open("rb") as file_obj:
        if file_name.endswith(".xlsx"):
            return _preview_xlsx(file_obj)
        if file_name.endswith(".xls"):
            return _preview_html_excel(file_obj)
    raise ValidationError("Solo se puede previsualizar archivos Excel .xls o .xlsx.")


def _append_spreadsheet_to_workbook(target_workbook, version, stage_order):
    from copy import copy

    file_name = version.file.name.lower()
    with version.file.open("rb") as file_obj:
        if file_name.endswith(".xlsx"):
            from openpyxl import load_workbook

            source_workbook = load_workbook(file_obj, data_only=False)
            for source_sheet in source_workbook.worksheets:
                desired_title = f"Etapa {stage_order}"
                if len(source_workbook.worksheets) > 1:
                    desired_title = f"E{stage_order} {source_sheet.title}"
                target_sheet = target_workbook.create_sheet(_unique_sheet_title(target_workbook, desired_title))
                target_sheet.freeze_panes = source_sheet.freeze_panes
                target_sheet.sheet_view.showGridLines = source_sheet.sheet_view.showGridLines
                target_sheet.auto_filter.ref = source_sheet.auto_filter.ref
                target_sheet.sheet_format = copy(source_sheet.sheet_format)
                target_sheet.sheet_properties = copy(source_sheet.sheet_properties)
                target_sheet.page_margins = copy(source_sheet.page_margins)
                target_sheet.page_setup = copy(source_sheet.page_setup)
                target_sheet.print_options = copy(source_sheet.print_options)
                for key, dimension in source_sheet.column_dimensions.items():
                    target_dimension = target_sheet.column_dimensions[key]
                    target_dimension.width = dimension.width
                    target_dimension.hidden = dimension.hidden
                    target_dimension.bestFit = dimension.bestFit
                    target_dimension.outlineLevel = dimension.outlineLevel
                for index, dimension in source_sheet.row_dimensions.items():
                    target_dimension = target_sheet.row_dimensions[index]
                    target_dimension.height = dimension.height
                    target_dimension.hidden = dimension.hidden
                    target_dimension.outlineLevel = dimension.outlineLevel
                for row in source_sheet.iter_rows():
                    for source_cell in row:
                        target_cell = target_sheet.cell(source_cell.row, source_cell.column, source_cell.value)
                        if source_cell.has_style:
                            target_cell.font = copy(source_cell.font)
                            target_cell.fill = copy(source_cell.fill)
                            target_cell.border = copy(source_cell.border)
                            target_cell.alignment = copy(source_cell.alignment)
                            target_cell.number_format = source_cell.number_format
                            target_cell.protection = copy(source_cell.protection)
                        if source_cell.hyperlink:
                            target_cell._hyperlink = copy(source_cell.hyperlink)
                        if source_cell.comment:
                            target_cell.comment = copy(source_cell.comment)
                for merged_range in source_sheet.merged_cells.ranges:
                    target_sheet.merge_cells(str(merged_range))
            sheet_count = len(source_workbook.worksheets)
            source_workbook.close()
            return sheet_count
        if file_name.endswith(".xls"):
            from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
            from openpyxl.utils import get_column_letter

            parser = _ExcelHtmlTableParser()
            parser.feed(file_obj.read().decode("utf-8-sig", errors="ignore"))
            if parser.rows:
                sheet = target_workbook.create_sheet(_unique_sheet_title(target_workbook, f"Etapa {stage_order}"))
                thin = Side(style="thin", color="000000")
                table_border = Border(left=thin, right=thin, top=thin, bottom=thin)
                start_row = 1
                max_columns = max(len(row) for row in parser.rows)
                if parser.caption:
                    sheet.cell(row=1, column=1, value=parser.caption)
                    sheet.cell(row=1, column=1).font = Font(bold=True, size=14)
                    sheet.cell(row=1, column=1).alignment = Alignment(horizontal="center")
                    if max_columns > 1:
                        sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max_columns)
                    start_row = 2
                column_widths = {}
                for row_offset, row in enumerate(parser.rows):
                    cell_types = parser.cell_types[row_offset]
                    for column_index, value in enumerate(row, start=1):
                        cell = sheet.cell(row=start_row + row_offset, column=column_index, value=value)
                        cell.border = table_border
                        cell.alignment = Alignment(vertical="top", wrap_text=True)
                        if cell_types[column_index - 1] == "th":
                            cell.font = Font(bold=True)
                            cell.fill = PatternFill("solid", fgColor="E7E9ED")
                        column_widths[column_index] = min(45, max(column_widths.get(column_index, 0), len(value) + 2))
                for column_index, width in column_widths.items():
                    sheet.column_dimensions[get_column_letter(column_index)].width = max(12, width)
                return 1
    raise ValidationError(f"No se pudo consolidar el archivo {version.file.name.rsplit('/', 1)[-1]} como Excel.")


def _append_spreadsheet_to_single_sheet(target_sheet, version, stage_order, start_row):
    from copy import copy

    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    def add_section_heading(title, row, column_count):
        target_sheet.cell(row=row, column=1, value=title)
        target_sheet.cell(row=row, column=1).font = Font(bold=True, color="FFFFFF", size=12)
        target_sheet.cell(row=row, column=1).fill = PatternFill("solid", fgColor="333333")
        target_sheet.cell(row=row, column=1).alignment = Alignment(vertical="center")
        if column_count > 1:
            target_sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=column_count)
        target_sheet.row_dimensions[row].height = 22

    file_name = version.file.name.lower()
    with version.file.open("rb") as file_obj:
        if file_name.endswith(".xlsx"):
            from openpyxl import load_workbook

            source_workbook = load_workbook(file_obj, data_only=False)
            for source_sheet in source_workbook.worksheets:
                column_count = max(1, source_sheet.max_column)
                heading = f"Etapa {stage_order} — {source_sheet.title}" if len(source_workbook.worksheets) > 1 else f"Etapa {stage_order}"
                add_section_heading(heading, start_row, column_count)
                content_start = start_row + 1
                for key, dimension in source_sheet.column_dimensions.items():
                    current_width = target_sheet.column_dimensions[key].width or 0
                    target_sheet.column_dimensions[key].width = max(current_width, dimension.width or 0)
                for index, dimension in source_sheet.row_dimensions.items():
                    target_sheet.row_dimensions[content_start + index - 1].height = dimension.height
                    target_sheet.row_dimensions[content_start + index - 1].hidden = dimension.hidden
                for row in source_sheet.iter_rows():
                    for source_cell in row:
                        target_cell = target_sheet.cell(content_start + source_cell.row - 1, source_cell.column, source_cell.value)
                        if source_cell.has_style:
                            target_cell.font = copy(source_cell.font)
                            target_cell.fill = copy(source_cell.fill)
                            target_cell.border = copy(source_cell.border)
                            target_cell.alignment = copy(source_cell.alignment)
                            target_cell.number_format = source_cell.number_format
                            target_cell.protection = copy(source_cell.protection)
                        if source_cell.hyperlink:
                            target_cell._hyperlink = copy(source_cell.hyperlink)
                        if source_cell.comment:
                            target_cell.comment = copy(source_cell.comment)
                for merged_range in source_sheet.merged_cells.ranges:
                    target_sheet.merge_cells(
                        start_row=content_start + merged_range.min_row - 1,
                        start_column=merged_range.min_col,
                        end_row=content_start + merged_range.max_row - 1,
                        end_column=merged_range.max_col,
                    )
                start_row = content_start + max(1, source_sheet.max_row) + 2
            source_workbook.close()
            return start_row
        if file_name.endswith(".xls"):
            parser = _ExcelHtmlTableParser()
            parser.feed(file_obj.read().decode("utf-8-sig", errors="ignore"))
            if parser.rows:
                column_count = max(len(row) for row in parser.rows)
                add_section_heading(f"Etapa {stage_order}", start_row, column_count)
                content_start = start_row + 1
                thin = Side(style="thin", color="000000")
                table_border = Border(left=thin, right=thin, top=thin, bottom=thin)
                if parser.caption:
                    target_sheet.cell(row=content_start, column=1, value=parser.caption)
                    target_sheet.cell(row=content_start, column=1).font = Font(bold=True, size=14)
                    target_sheet.cell(row=content_start, column=1).alignment = Alignment(horizontal="center")
                    if column_count > 1:
                        target_sheet.merge_cells(start_row=content_start, start_column=1, end_row=content_start, end_column=column_count)
                    content_start += 1
                for row_offset, row in enumerate(parser.rows):
                    cell_types = parser.cell_types[row_offset]
                    for column_index, value in enumerate(row, start=1):
                        cell = target_sheet.cell(row=content_start + row_offset, column=column_index, value=value)
                        cell.border = table_border
                        cell.alignment = Alignment(vertical="top", wrap_text=True)
                        if cell_types[column_index - 1] == "th":
                            cell.font = Font(bold=True)
                            cell.fill = PatternFill("solid", fgColor="E7E9ED")
                        current_width = target_sheet.column_dimensions[cell.column_letter].width or 0
                        target_sheet.column_dimensions[cell.column_letter].width = max(current_width, min(45, max(12, len(value) + 2)))
                return content_start + len(parser.rows) + 2
    raise ValidationError(f"No se pudo consolidar el archivo {version.file.name.rsplit('/', 1)[-1]} como Excel.")


def _unique_sheet_title(workbook, desired_title):
    invalid_characters = set("[]:*?/\\")
    base = "".join("_" if character in invalid_characters else character for character in desired_title).strip()[:31] or "Hoja"
    title = base
    suffix = 2
    while title in workbook.sheetnames:
        marker = f"_{suffix}"
        title = f"{base[:31 - len(marker)]}{marker}"
        suffix += 1
    return title


def _latest_available_version(document):
    for version in document.versions.order_by("-version_number"):
        if version.file and version.file.storage.exists(version.file.name):
            return version
    return None


# Construye un evento de calendario a partir de etapa, documento, visita o evento manual.
def _calendar_section(section):
    if not section:
        return None
    return {
        "id": section.id,
        "name": section.name,
        "course": section.course.name,
        "campus": section.campus_career.campus.name,
        "career": section.campus_career.career.name,
    }


def _calendar_event(event_id, event_type, title, event_date, section, status="", metadata=None):
    return {
        "id": event_id,
        "type": event_type,
        "title": title,
        "date": event_date.isoformat() if event_date else None,
        "status": status,
        "section": _calendar_section(section),
        "metadata": metadata or {},
    }


def _custom_calendar_event_to_dict(event, user):
    return _calendar_event(
        f"custom-{event.id}",
        "custom",
        event.title,
        event.event_date,
        event.section,
        event.category,
        {
            "event_id": event.id,
            "category": event.category,
            "start_time": event.start_time.isoformat(timespec="minutes") if event.start_time else "",
            "end_time": event.end_time.isoformat(timespec="minutes") if event.end_time else "",
            "location": event.location,
            "description": event.description,
            "created_by": (event.created_by.get_full_name() or event.created_by.username) if event.created_by else "Sistema",
            "editable": event.created_by_id == user.id or has_role(user, ADMIN_ROLES),
        },
    )


# Entrega una cookie CSRF para clientes que mantienen sesión Django.
@ensure_csrf_cookie
@require_GET
@never_cache
def csrf_token(request):
    return JsonResponse({"detail": "CSRF cookie set", "csrfToken": get_token(request)})


# Inicia sesión usando el sistema de autenticacion de Django.
@require_POST
@never_cache
def login_view(request):
    payload = _json_body(request)
    if payload is None:
        return HttpResponseBadRequest("JSON inválido.")
    username = payload.get("username") or payload.get("email")
    password = payload.get("password")
    if not username or not password:
        return HttpResponseBadRequest("Usuario y contraseña son obligatorios.")
    if "@" in username:
        user_by_email = User.objects.filter(email__iexact=username).first()
        username = user_by_email.username if user_by_email else username
    user = authenticate(request, username=username, password=password)
    if user is None:
        return JsonResponse({"detail": "Los datos ingresados son incorrectos o la cuenta está inhabilitada."}, status=400)
    login(request, user)
    return JsonResponse({"user": _session_user(user)})


# Cierra la sesión Django activa.
@login_required
@require_POST
@never_cache
def logout_view(request):
    logout(request)
    return JsonResponse({"detail": "Sesion cerrada."})


# Devuelve el usuario actual para restaurar sesión en React.
@login_required
@require_GET
@never_cache
def me_view(request):
    return JsonResponse(_session_user(request.user))


def _session_user(user):
    roles = sorted(set(user.groups.filter(name__in=User.Role.values).values_list("name", flat=True)) | {user.rol})
    return {"id": user.id, "username": user.username, "email": user.email, "rol": user.rol, "roles": roles}


# Endpoint simple para verificar que el servicio responde.
@require_GET
def health_check(request):
    return JsonResponse({"status": "ok", "service": "ConectaAPI"})


@login_required
@require_GET
@never_cache
def notification_list(request):
    generate_due_alerts()
    alerts = Alert.objects.filter(recipients=request.user).select_related("section__course", "stage").order_by("-created_at")[:50]
    read_ids = set(AlertRead.objects.filter(user=request.user, alert__in=alerts).values_list("alert_id", flat=True))
    results = [{
        "id": alert.id,
        "title": alert.title,
        "message": alert.message,
        "severity": alert.severity,
        "created_at": alert.created_at.isoformat(),
        "is_read": alert.id in read_ids,
        "section": f"{alert.section.course.name} · {alert.section.name}",
        "url": f"#/sections/{alert.section_id}?stage={alert.stage.order}" if alert.stage else f"#/sections/{alert.section_id}",
    } for alert in alerts]
    return JsonResponse({"unread_count": sum(not item["is_read"] for item in results), "results": results})


@login_required
@require_POST
def notification_read(request, alert_id):
    alert = get_object_or_404(Alert, pk=alert_id, recipients=request.user)
    AlertRead.objects.get_or_create(alert=alert, user=request.user)
    return JsonResponse({"detail": "Alerta marcada como leída."})


ADMIN_RESOURCES = {
    "users": (User, ("username", "email", "first_name", "last_name", "rol", "phone", "is_active")),
    "campuses": (Campus, ("name", "is_active")),
    "careers": (Career, ("code", "name", "is_active")),
    "courses": (Course, ("code", "name", "is_active")),
    "semesters": (Semester, ("year", "period", "is_active")),
    "partners": (ProductivePartner, ("name", "rut", "business_area", "contact_name", "contact_email", "contact_phone", "address", "requirements_match", "status")),
    "sections": (APISection, ("name", "campus_career_id", "course_id", "semester_id", "career_director_id", "specialist_teacher_id", "ctp_coordinator_id", "productive_partner_id", "status")),
    "documents": (Document, ("stage_id", "title", "document_type", "status", "delivery_due_date", "review_due_date", "is_historical")),
    "document-versions": (DocumentVersion, ("notes",)),
    "campus-careers": (CampusCareer, ("campus_id", "career_id", "is_active")),
    "user-campuses": (UserCampus, ("user_id", "campus_id", "is_active")),
    "stages": (ProcessStage, ("section_id", "stage_type", "order", "status", "due_date", "completed_at", "notes")),
    "visits": (Visit, ("section_id", "visit_number", "scheduled_date", "place", "assigned_teacher_id", "status", "notes")),
    "admin-alerts": (Alert, ("section_id", "stage_id", "alert_type", "severity", "title", "message", "due_date", "is_active", "email_sent", "resolved_at")),
    "audit-logs": (AuditLog, ()),
    "ctp-reviews": (CTPValidationReview, ()),
}

SECTION_IMPORT_HEADERS = (
    "Programa de Estudio", "Cód Asignatura", "Asignatura", "Sección",
    "Rut Prof.", "Profesor", "Director de Carrera", "Correo DC",
    "CTP", "Correo CTP", "Docente Especialista", "Correo DE",
)


def _section_import_user(name, email, role, campus):
    name = str(name).strip()
    email = str(email).strip().lower()
    user = User.objects.filter(email__iexact=email).first()
    created = user is None
    if created:
        username = name
        suffix = 2
        while User.objects.filter(username=username).exists():
            username = f"{name} {suffix}"
            suffix += 1
        user = User(username=username, email=email, rol=role, is_active=True)
        user.set_unusable_password()
        user.save()
    else:
        fields = []
        if not user.is_active:
            user.is_active = True
            fields.append("is_active")
        if user.rol != role:
            user.rol = role
            fields.append("rol")
        if fields:
            user.save(update_fields=fields)
    group, _ = Group.objects.get_or_create(name=role)
    user.groups.add(group)
    UserCampus.objects.update_or_create(user=user, campus=campus, defaults={"is_active": True})
    return user, created


@login_required
@require_POST
def section_bulk_import(request):
    if not _admin_only(request):
        return _forbidden()
    uploaded = request.FILES.get("file")
    if not uploaded:
        return _validation_error(ValidationError("Selecciona un archivo Excel."))
    if not uploaded.name.lower().endswith(".xlsx"):
        return _validation_error(ValidationError("La importación requiere un archivo .xlsx."))
    if uploaded.size > 10 * 1024 * 1024:
        return _validation_error(ValidationError("El archivo supera el máximo de 10 MB."))
    try:
        campus = Campus.objects.get(pk=request.POST.get("campus_id"), is_active=True)
        semester = Semester.objects.get(pk=request.POST.get("semester_id"), is_active=True)
    except (Campus.DoesNotExist, Semester.DoesNotExist, ValueError, TypeError):
        return _validation_error(ValidationError("Selecciona una sede y un semestre activos."))

    try:
        from openpyxl import load_workbook
        workbook = load_workbook(uploaded, read_only=True, data_only=True)
        sheet = workbook.active
        headers = tuple(cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1)))
        if headers != SECTION_IMPORT_HEADERS:
            raise ValidationError("Las 12 columnas del Excel no corresponden al formato requerido.")
        rows = []
        required = (
            "Programa de Estudio", "Cód Asignatura", "Asignatura", "Sección",
            "Director de Carrera", "Correo DC", "CTP", "Correo CTP",
            "Docente Especialista", "Correo DE",
        )
        for row_number, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
            if not any(values):
                continue
            row = dict(zip(SECTION_IMPORT_HEADERS, values))
            missing = [field for field in required if not str(row.get(field) or "").strip()]
            if missing:
                raise ValidationError(f"Fila {row_number}: faltan {', '.join(missing)}.")
            program = [part.strip() for part in str(row["Programa de Estudio"]).split(" - ", 1)]
            if len(program) != 2 or not all(program):
                raise ValidationError(f"Fila {row_number}: Programa de Estudio debe usar 'CÓDIGO - Nombre'.")
            for field in ("Correo DC", "Correo CTP", "Correo DE"):
                try:
                    validate_email(str(row[field]).strip())
                except ValidationError as error:
                    raise ValidationError(f"Fila {row_number}: {field} no es un correo válido.") from error
            row["career_code"], row["career_name"] = program
            rows.append(row)
        if not rows:
            raise ValidationError("El archivo no contiene secciones para importar.")
    except ValidationError as error:
        return _validation_error(error)
    except Exception:
        return _validation_error(ValidationError("No se pudo leer el Excel. Verifica que el archivo no esté dañado."))

    created_sections = updated_sections = created_users = 0
    try:
        with transaction.atomic():
            for row in rows:
                career, _ = Career.objects.update_or_create(
                    code=row["career_code"], defaults={"name": row["career_name"], "is_active": True}
                )
                campus_career, _ = CampusCareer.objects.update_or_create(
                    campus=campus, career=career, defaults={"is_active": True}
                )
                course, _ = Course.objects.update_or_create(
                    code=str(row["Cód Asignatura"]).strip(),
                    defaults={"name": str(row["Asignatura"]).strip(), "is_active": True},
                )
                director, created = _section_import_user(row["Director de Carrera"], row["Correo DC"], User.Role.DC, campus)
                created_users += int(created)
                ctp, created = _section_import_user(row["CTP"], row["Correo CTP"], User.Role.CTP, campus)
                created_users += int(created)
                specialist, created = _section_import_user(row["Docente Especialista"], row["Correo DE"], User.Role.DE, campus)
                created_users += int(created)
                _, created = APISection.objects.update_or_create(
                    campus_career=campus_career,
                    course=course,
                    semester=semester,
                    name=str(row["Sección"]).strip(),
                    defaults={
                        "career_director": director,
                        "ctp_coordinator": ctp,
                        "specialist_teacher": specialist,
                    },
                )
                created_sections += int(created)
                updated_sections += int(not created)
    except (ValidationError, IntegrityError, ValueError, TypeError) as error:
        return _validation_error(error)

    return JsonResponse({
        "detail": "Importación masiva completada.",
        "created_sections": created_sections,
        "updated_sections": updated_sections,
        "created_users": created_users,
        "processed_rows": len(rows),
    })


@login_required
@require_GET
def section_bulk_export(request):
    if not _admin_only(request):
        return _forbidden()
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Hoja1"
    sheet.append(SECTION_IMPORT_HEADERS)
    sections = APISection.objects.select_related(
        "campus_career__career", "course", "career_director", "ctp_coordinator", "specialist_teacher"
    ).order_by("campus_career__career__code", "course__code", "name")
    for section in sections:
        career = section.campus_career.career
        sheet.append((
            f"{career.code} - {career.name}", section.course.code, section.course.name, section.name,
            "", "",
            section.career_director.get_full_name() or section.career_director.username if section.career_director else "",
            section.career_director.email if section.career_director else "",
            section.ctp_coordinator.get_full_name() or section.ctp_coordinator.username if section.ctp_coordinator else "",
            section.ctp_coordinator.email if section.ctp_coordinator else "",
            section.specialist_teacher.get_full_name() or section.specialist_teacher.username if section.specialist_teacher else "",
            section.specialist_teacher.email if section.specialist_teacher else "",
        ))
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="0E4684")
    for column in sheet.columns:
        letter = column[0].column_letter
        sheet.column_dimensions[letter].width = min(45, max(14, max(len(str(cell.value or "")) for cell in column) + 2))
    output = io.BytesIO()
    workbook.save(output)
    response = HttpResponse(output.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = 'attachment; filename="secciones-api.xlsx"'
    return response

ADMIN_RESOURCE_LABELS = {
    "users": "Usuarios",
    "campuses": "Sedes",
    "careers": "Carreras",
    "courses": "Asignaturas API",
    "semesters": "Semestres académicos",
    "partners": "Socios productivos",
    "sections": "Secciones API",
    "documents": "Documentos de etapas",
    "document-versions": "Versiones de documentos",
    "campus-careers": "Carreras por sede",
    "user-campuses": "Sedes asignadas a usuarios",
    "stages": "Etapas del proceso",
    "visits": "Visitas de seguimiento",
    "admin-alerts": "Alertas operativas",
    "ctp-reviews": "Revisiones CTP",
}


def _admin_record(resource, obj):
    data = {"id": obj.id}
    for field in ADMIN_RESOURCES[resource][1]:
        data[field] = getattr(obj, field)
    if resource == "users":
        data["campus_ids"] = list(obj.campus_access.filter(is_active=True).values_list("campus_id", flat=True))
        data["rol_display"] = obj.get_rol_display()
        data["roles"] = sorted(set(obj.groups.filter(name__in=User.Role.values).values_list("name", flat=True)) | {obj.rol})
    if resource == "sections":
        data["campus_career_label"] = f"{obj.campus_career.campus.name} · {obj.campus_career.career.name}"
        data["course_label"] = f"{obj.course.code} · {obj.course.name}"
        data["semester_label"] = str(obj.semester)
        data["career_director_label"] = (obj.career_director.get_full_name() or obj.career_director.username) if obj.career_director else "Sin asignar"
        data["specialist_teacher_id"] = obj.specialist_teacher_id
        data["ctp_coordinator_id"] = obj.ctp_coordinator_id
        data["specialist_label"] = (obj.specialist_teacher.get_full_name() or obj.specialist_teacher.username) if obj.specialist_teacher else "Sin asignar"
        data["ctp_user_label"] = (obj.ctp_coordinator.get_full_name() or obj.ctp_coordinator.username) if obj.ctp_coordinator else "Sin asignar"
    if resource == "documents":
        data["stage_label"] = f"{obj.stage.section.course.name} · {obj.stage.section.name} · {obj.stage.stage_type}"
        data["current_version"] = obj.current_version
        data["version_count"] = obj.versions.count()
    if resource == "document-versions":
        data.update({
            "document_id": obj.document_id,
            "document_label": obj.document.title,
            "section_label": f"{obj.document.stage.section.course.name} · {obj.document.stage.section.name}",
            "version_number": obj.version_number,
            "file_name": obj.file.name.rsplit("/", 1)[-1],
            "uploaded_by": obj.uploaded_by.get_full_name() or obj.uploaded_by.username if obj.uploaded_by else "Sistema",
            "created_at": obj.created_at.isoformat(),
        })
    if resource == "campus-careers":
        data.update({"name": f"{obj.campus.name} · {obj.career.name}", "campus_label": obj.campus.name, "career_label": obj.career.name})
    if resource == "user-campuses":
        data.update({"name": obj.user.get_full_name() or obj.user.username, "user_label": obj.user.get_full_name() or obj.user.username, "campus_label": obj.campus.name})
    if resource == "stages":
        data.update({"name": obj.get_stage_type_display(), "section_label": f"{obj.section.course.name} · {obj.section.name}", "responsible_user_label": (obj.responsible_user.get_full_name() or obj.responsible_user.username) if obj.responsible_user else "Sin asignar"})
    if resource == "visits":
        data.update({"name": f"Visita {obj.visit_number}", "section_label": f"{obj.section.course.name} · {obj.section.name}", "assigned_teacher_label": (obj.assigned_teacher.get_full_name() or obj.assigned_teacher.username) if obj.assigned_teacher else "Sin docente"})
    if resource == "admin-alerts":
        data.update({"name": obj.title, "section_label": f"{obj.section.course.name} · {obj.section.name}", "recipient_ids": list(obj.recipients.values_list("id", flat=True))})
    if resource == "audit-logs":
        data.update({"name": obj.get_action_display(), "action": obj.action, "actor_label": (obj.actor.get_full_name() or obj.actor.username) if obj.actor else "Sistema", "section_label": str(obj.section) if obj.section else "Sin sección", "created_at": obj.created_at.isoformat(), "metadata": obj.metadata, "ip_address": obj.ip_address})
    if resource == "ctp-reviews":
        data.update({"name": f"Revisión CTP {obj.attempt}", "section_label": f"{obj.section.course.name} · {obj.section.name}", "result_display": obj.get_result_display(), "submitted_by_label": (obj.submitted_by.get_full_name() or obj.submitted_by.username) if obj.submitted_by else "Sistema", "reviewed_by_label": (obj.reviewed_by.get_full_name() or obj.reviewed_by.username) if obj.reviewed_by else "Pendiente", "submitted_at": obj.submitted_at.isoformat(), "reviewed_at": obj.reviewed_at.isoformat() if obj.reviewed_at else None, "feedback": obj.feedback, "error_stages": obj.error_stages, "details": obj.details})
    return data


def _audit_object_label(obj):
    if isinstance(obj, User):
        return obj.get_full_name() or obj.username
    if isinstance(obj, DocumentVersion):
        return f"{obj.document.title} · v{obj.version_number}"
    return str(obj)


def _audit_related(obj):
    section = None
    stage = None
    document = None
    if isinstance(obj, APISection):
        section = obj
    elif isinstance(obj, ProcessStage):
        section = obj.section
        stage = obj
    elif isinstance(obj, Document):
        section = obj.stage.section
        stage = obj.stage
        document = obj
    elif isinstance(obj, DocumentVersion):
        section = obj.document.stage.section
        stage = obj.document.stage
        document = obj.document
    elif isinstance(obj, Visit):
        section = obj.section
    elif isinstance(obj, Alert):
        section = obj.section
        stage = obj.stage
    elif isinstance(obj, CTPValidationReview):
        section = obj.section
    elif isinstance(obj, CalendarEvent):
        section = obj.section
    return section, stage, document


def _audit_admin_action(request, action, resource, obj, metadata=None, include_related=True):
    section, stage, document = _audit_related(obj) if include_related else (None, None, None)
    payload = {
        "resource": resource,
        "resource_label": ADMIN_RESOURCE_LABELS.get(resource, resource),
        "object_id": obj.pk,
        "object_label": _audit_object_label(obj),
    }
    payload.update(metadata or {})
    log_action(action, actor=request.user, section=section, stage=stage, document=document, metadata=payload, ip_address=_client_ip(request))


def _audit_calendar_action(request, action, event, metadata=None, include_related=True):
    section, stage, document = _audit_related(event) if include_related else (None, None, None)
    payload = {
        "event_id": event.pk,
        "event_title": getattr(event, "title", str(event)),
        "event_date": event.event_date.isoformat() if getattr(event, "event_date", None) else None,
        "category": getattr(event, "category", ""),
    }
    payload.update(metadata or {})
    log_action(action, actor=request.user, section=section, stage=stage, document=document, metadata=payload, ip_address=_client_ip(request))


def _admin_only(request):
    return has_role(request.user, {User.Role.ADMIN})


def user_option(item):
    return {
        "value": item.id,
        "label": item.get_full_name() or item.username,
        "username": item.username,
        "email": item.email,
        "first_name": item.first_name,
        "last_name": item.last_name,
    }


@login_required
@require_GET
def administration_counts(request):
    if not _admin_only(request):
        return _forbidden()
    return JsonResponse({resource: model.objects.count() for resource, (model, _fields) in ADMIN_RESOURCES.items()})

@login_required
@require_GET
def administration_options(request):
    if not _admin_only(request):
        return _forbidden()
    return JsonResponse({
        "roles": [{"value": value, "label": label} for value, label in User.Role.choices],
        "periods": [{"value": value, "label": label} for value, label in Semester.Period.choices],
        "partner_statuses": [{"value": value, "label": label} for value, label in ProductivePartner.Status.choices],
        "section_statuses": [{"value": value, "label": label} for value, label in APISection.Status.choices],
        "document_types": [{"value": value, "label": label} for value, label in Document.DocumentType.choices],
        "document_statuses": [{"value": value, "label": label} for value, label in Document.Status.choices],
        "campuses": [{"id": campus.id, "name": campus.name} for campus in Campus.objects.filter(is_active=True).order_by("name")],
        "campus_careers": [{"value": item.id, "label": f"{item.campus.name} · {item.career.name}"} for item in CampusCareer.objects.select_related("campus", "career").filter(is_active=True, campus__is_active=True, career__is_active=True).order_by("campus__name", "career__name")],
        "courses": [{"value": item.id, "label": f"{item.code} · {item.name}"} for item in Course.objects.filter(is_active=True).order_by("name")],
        "semesters": [{"value": item.id, "label": str(item)} for item in Semester.objects.filter(is_active=True).order_by("-year", "period")],
        "directors": [user_option(item) for item in User.objects.filter(Q(rol=User.Role.DC) | Q(groups__name=User.Role.DC), is_active=True).distinct().order_by("first_name", "last_name", "username")],
        "specialists": [user_option(item) for item in User.objects.filter(Q(rol=User.Role.DE) | Q(groups__name=User.Role.DE), is_active=True).distinct().order_by("first_name", "last_name", "username")],
        "ctps": [user_option(item) for item in User.objects.filter(Q(rol=User.Role.CTP) | Q(groups__name=User.Role.CTP), is_active=True).distinct().order_by("first_name", "last_name", "username")],
        "partners": [{"value": item.id, "label": item.name} for item in ProductivePartner.objects.exclude(status=ProductivePartner.Status.INACTIVE).order_by("name")],
        "stages": [{"value": item.id, "label": f"{item.section.course.name} · {item.section.name} · {item.stage_type}"} for item in ProcessStage.objects.select_related("section__course", "section__campus_career__campus", "section__campus_career__career").filter(section__status__in=[APISection.Status.ACTIVE, APISection.Status.IN_PROGRESS], section__course__is_active=True, section__semester__is_active=True, section__campus_career__is_active=True, section__campus_career__campus__is_active=True, section__campus_career__career__is_active=True).order_by("section", "order")],
        "stage_types": [{"value": value, "label": label} for value, label in ProcessStage.StageType.choices],
        "stage_statuses": [{"value": value, "label": label} for value, label in ProcessStage.Status.choices],
        "visit_statuses": [{"value": value, "label": label} for value, label in Visit.Status.choices],
        "alert_types": [{"value": value, "label": label} for value, label in Alert.AlertType.choices],
        "alert_severities": [{"value": value, "label": label} for value, label in Alert.Severity.choices],
        "careers": [{"value": item.id, "label": f"{item.code} · {item.name}"} for item in Career.objects.filter(is_active=True).order_by("name")],
        "sections": [{"value": item.id, "label": f"{item.course.name} · {item.name}"} for item in APISection.objects.select_related("course", "campus_career__campus", "campus_career__career").filter(status__in=[APISection.Status.ACTIVE, APISection.Status.IN_PROGRESS], course__is_active=True, semester__is_active=True, campus_career__is_active=True, campus_career__campus__is_active=True, campus_career__career__is_active=True).order_by("course__name", "name")],
        "users": [user_option(item) for item in User.objects.filter(is_active=True).order_by("first_name", "last_name", "username")],
    })


def _template_record(template):
    uploader = template.uploaded_by
    return {
        "id": template.id,
        "stage_type": template.stage_type,
        "stage_label": template.get_stage_type_display(),
        "file_name": template.file.name.rsplit("/", 1)[-1],
        "notes": template.notes,
        "uploaded_by": (uploader.get_full_name() or uploader.username) if uploader else "Sistema",
        "updated_at": template.updated_at.isoformat(),
    }


@login_required
@require_http_methods(["GET", "POST"])
def cycle_template_collection(request):
    if not _admin_only(request):
        return _forbidden()
    if request.method == "GET":
        templates = CycleTemplate.objects.select_related("uploaded_by")
        return JsonResponse({"count": templates.count(), "results": [_template_record(item) for item in templates]})
    try:
        template = CycleTemplate(stage_type=request.POST.get("stage_type", ""), file=request.FILES.get("file"), notes=request.POST.get("notes", ""), uploaded_by=request.user)
        template.full_clean()
        template.save()
        return JsonResponse(_template_record(template), status=201)
    except (ValidationError, IntegrityError) as error:
        return _validation_error(error)


@login_required
@require_http_methods(["POST", "DELETE"])
def cycle_template_detail(request, template_id):
    if not _admin_only(request):
        return _forbidden()
    template = get_object_or_404(CycleTemplate, pk=template_id)
    if request.method == "DELETE":
        template.file.delete(save=False)
        template.delete()
        return JsonResponse({"detail": "Plantilla eliminada."})
    old_file = template.file
    template.stage_type = request.POST.get("stage_type", template.stage_type)
    template.notes = request.POST.get("notes", template.notes)
    new_file = request.FILES.get("file")
    if new_file:
        template.file = new_file
    template.uploaded_by = request.user
    try:
        template.full_clean()
        template.save()
        if new_file and old_file.name != template.file.name:
            old_file.delete(save=False)
        return JsonResponse(_template_record(template))
    except (ValidationError, IntegrityError) as error:
        return _validation_error(error)


@login_required
@require_GET
def cycle_template_download(request, template_id):
    template = get_object_or_404(CycleTemplate, pk=template_id)
    return FileResponse(template.file.open("rb"), as_attachment=True, filename=template.file.name.rsplit("/", 1)[-1])


@login_required
@require_http_methods(["GET", "POST"])
def administration_collection(request, resource):
    if not _admin_only(request):
        return _forbidden()
    config = ADMIN_RESOURCES.get(resource)
    if config is None:
        return JsonResponse({"detail": "Recurso administrativo desconocido."}, status=404)
    model, fields = config
    if request.method == "GET":
        queryset = model.objects.all().order_by("id")
        if resource == "users":
            queryset = queryset.prefetch_related("groups", "campus_access")
        elif resource == "sections":
            queryset = queryset.select_related("campus_career__campus", "campus_career__career", "course", "semester", "career_director", "specialist_teacher", "ctp_coordinator")
        elif resource == "documents":
            queryset = queryset.select_related("stage__section__course").prefetch_related("versions")
        elif resource == "document-versions":
            queryset = queryset.select_related("document__stage__section__course", "uploaded_by")
        elif resource == "campus-careers":
            queryset = queryset.select_related("campus", "career")
        elif resource == "user-campuses":
            queryset = queryset.select_related("user", "campus")
        elif resource == "stages":
            queryset = queryset.select_related("section__course", "section__career_director", "section__specialist_teacher", "section__ctp_coordinator")
        elif resource == "visits":
            queryset = queryset.select_related("section__course", "assigned_teacher")
        elif resource == "admin-alerts":
            queryset = queryset.select_related("section__course", "stage").prefetch_related("recipients")
        elif resource == "audit-logs":
            queryset = queryset.select_related("actor", "section").order_by("-created_at", "-id")
        elif resource == "ctp-reviews":
            queryset = queryset.select_related("section__course", "submitted_by", "reviewed_by")
        total_count = queryset.count()
        if resource in {"users", "campuses", "careers", "campus-careers", "user-campuses", "courses", "sections", "semesters", "partners", "documents", "document-versions", "stages", "visits", "admin-alerts", "audit-logs", "ctp-reviews"}:
            try:
                page = max(1, int(request.GET.get("page", 1)))
                page_size = min(10, max(1, int(request.GET.get("page_size", 10))))
            except (TypeError, ValueError):
                return JsonResponse({"detail": "Parámetros de paginación inválidos."}, status=400)
            total_pages = max(1, (total_count + page_size - 1) // page_size)
            page = min(page, total_pages)
            start = (page - 1) * page_size
            queryset = queryset[start:start + page_size]
            return JsonResponse({
                "count": total_count,
                "page": page,
                "page_size": page_size,
                "total_pages": total_pages,
                "results": [_admin_record(resource, obj) for obj in queryset],
            })
        return JsonResponse({"count": total_count, "results": [_admin_record(resource, obj) for obj in queryset]})
    if resource in {"document-versions", "audit-logs", "ctp-reviews"}:
        return JsonResponse({"detail": "Las versiones se crean al subir archivos desde el flujo de Documentos."}, status=405)
    payload = _json_body(request)
    if payload is None:
        return HttpResponseBadRequest("JSON inválido.")
    try:
        values = {field: payload[field] for field in fields if field in payload}
        if resource == "users":
            roles = payload.get("roles", [payload.get("rol", User.Role.C_API)])
            if not roles:
                raise ValidationError("Selecciona al menos un rol.")
            values["rol"] = roles[0]
            password = payload.get("password")
            if not password:
                raise ValidationError("La contraseña es obligatoria al crear un usuario.")
            obj = User(**values)
            obj.set_password(password)
        else:
            obj = model(**values)
        obj.full_clean()
        _validate_administration_active_references(resource, obj)
        if resource == "users":
            _validate_user_campus_ids(obj, payload.get("campus_ids", []))
        obj.save()
        if resource == "users":
            _set_user_campuses(obj, payload.get("campus_ids", []))
            _set_user_roles(obj, roles)
        if resource == "sections":
            _set_section_stage_responsibles(obj, payload)
        if resource == "admin-alerts":
            _set_alert_recipients(obj, payload.get("recipient_ids", []))
        _audit_admin_action(request, AuditLog.Action.ADMIN_CREATED, resource, obj, {"changed_fields": sorted(values.keys())})
        return JsonResponse(_admin_record(resource, obj), status=201)
    except (ValidationError, IntegrityError, ValueError, TypeError) as error:
        return _validation_error(error)




def _validate_role_user(user_id, role, label):
    if not user_id:
        return None
    user = User.objects.filter(Q(rol=role) | Q(groups__name=role), pk=user_id, is_active=True).distinct().first()
    if not user:
        raise ValidationError(f"Selecciona un usuario activo con rol {label}.")
    return user


def _set_section_stage_responsibles(section, payload):
    assignments = [
        ("specialist_teacher_id", User.Role.DE, "Docente Especialista"),
        ("ctp_coordinator_id", User.Role.CTP, "Coordinador Técnico Pedagógico"),
        ("career_director_id", User.Role.DC, "Director de Carrera"),
    ]
    changed_fields = []
    for field, role, label in assignments:
        if field not in payload:
            continue
        responsible = _validate_role_user(payload.get(field), role, label)
        setattr(section, field, responsible.pk if responsible else None)
        changed_fields.append(field)
    if changed_fields:
        section.full_clean()
        section.save(update_fields=[*changed_fields, "updated_at"])

def _is_active_reference(obj):
    if obj is None:
        return True
    if hasattr(obj, "is_active") and not obj.is_active:
        return False
    if isinstance(obj, ProductivePartner) and obj.status == ProductivePartner.Status.INACTIVE:
        return False
    if isinstance(obj, APISection) and obj.status == APISection.Status.SUSPENDED:
        return False
    if isinstance(obj, CampusCareer):
        return obj.is_active and obj.campus.is_active and obj.career.is_active
    return True


def _ensure_active_reference(obj, label):
    if not _is_active_reference(obj):
        raise ValidationError(f"{label} está inactivo o inhabilitado y no se puede utilizar.")


def _validate_administration_active_references(resource, obj):
    if resource == "campus-careers":
        _ensure_active_reference(obj.campus, "La sede")
        _ensure_active_reference(obj.career, "La carrera")
    elif resource == "user-campuses":
        _ensure_active_reference(obj.user, "El usuario")
        _ensure_active_reference(obj.campus, "La sede")
    elif resource == "sections":
        _ensure_active_reference(obj.campus_career, "La relación entre sede y carrera")
        _ensure_active_reference(obj.course, "La asignatura API")
        _ensure_active_reference(obj.semester, "El semestre académico")
        _ensure_active_reference(obj.career_director, "El Director de Carrera")
        _ensure_active_reference(obj.specialist_teacher, "El Docente Especialista")
        _ensure_active_reference(obj.ctp_coordinator, "El Coordinador CTP")
        _ensure_active_reference(obj.productive_partner, "El socio productivo")
    elif resource == "documents":
        _ensure_active_reference(obj.stage.section, "La sección API")
    elif resource == "stages":
        _ensure_active_reference(obj.section, "La sección API")
    elif resource == "visits":
        _ensure_active_reference(obj.section, "La sección API")
        _ensure_active_reference(obj.assigned_teacher, "El docente asignado")
    elif resource == "admin-alerts":
        _ensure_active_reference(obj.section, "La sección API")
        _ensure_active_reference(obj.stage.section if obj.stage else None, "La etapa relacionada")


def _validate_user_campus_ids(user, campus_ids):
    selected_ids = set(campus_ids)
    if selected_ids and not user.is_active:
        raise ValidationError("El usuario está inhabilitado y no se le pueden asignar sedes.")
    valid_ids = set(Campus.objects.filter(id__in=selected_ids, is_active=True).values_list("id", flat=True))
    if len(valid_ids) != len(selected_ids):
        raise ValidationError("Una o más sedes seleccionadas no existen o están inactivas.")
    return valid_ids

def _set_user_campuses(user, campus_ids):
    valid_ids = _validate_user_campus_ids(user, campus_ids)
    UserCampus.objects.filter(user=user).exclude(campus_id__in=valid_ids).delete()
    for campus_id in valid_ids:
        access, _ = UserCampus.objects.get_or_create(user=user, campus_id=campus_id)
        if not access.is_active:
            access.is_active = True
            access.save(update_fields=["is_active", "updated_at"])


def _set_user_roles(user, roles):
    selected = set(roles)
    valid_roles = set(User.Role.values)
    if not selected or not selected.issubset(valid_roles):
        raise ValidationError("Uno o más roles seleccionados no son válidos.")
    user.groups.remove(*Group.objects.filter(name__in=valid_roles))
    user.groups.add(*(Group.objects.get_or_create(name=role)[0] for role in selected))
    if user.rol not in selected:
        user.rol = roles[0]
        user.save(update_fields=["rol"])


def _set_alert_recipients(alert, recipient_ids):
    recipients = User.objects.filter(id__in=recipient_ids, is_active=True)
    if recipients.count() != len(set(recipient_ids)):
        raise ValidationError("Uno o más destinatarios no existen o están inhabilitados.")
    alert.recipients.set(recipients)


@login_required
@require_http_methods(["PATCH", "DELETE"])
def administration_detail(request, resource, object_id):
    if not _admin_only(request):
        return _forbidden()
    config = ADMIN_RESOURCES.get(resource)
    if config is None:
        return JsonResponse({"detail": "Recurso administrativo desconocido."}, status=404)
    model, fields = config
    obj = get_object_or_404(model, pk=object_id)
    if resource in {"audit-logs", "ctp-reviews"}:
        return JsonResponse({"detail": "La auditoría es de solo lectura."}, status=405)
    if request.method == "DELETE":
        action = request.GET.get("action", "toggle")
        if action != "delete":
            if resource == "sections":
                if obj.status == APISection.Status.SUSPENDED:
                    obj.status = APISection.Status.ACTIVE
                    detail = "Sección API activada."
                else:
                    obj.status = APISection.Status.SUSPENDED
                    detail = "Sección API desactivada."
                obj.save(update_fields=["status", "updated_at"])
                _audit_admin_action(request, AuditLog.Action.ADMIN_TOGGLED, resource, obj, {"status": obj.status})
                return JsonResponse({"detail": detail})
            if resource == "partners":
                if obj.status == ProductivePartner.Status.INACTIVE:
                    obj.status = ProductivePartner.Status.PROPOSED
                    detail = "Socio productivo activado."
                else:
                    obj.status = ProductivePartner.Status.INACTIVE
                    detail = "Socio productivo desactivado."
                obj.save(update_fields=["status", "updated_at"])
                _audit_admin_action(request, AuditLog.Action.ADMIN_TOGGLED, resource, obj, {"status": obj.status})
                return JsonResponse({"detail": detail})
            if "is_active" in fields:
                if resource == "users" and obj.pk == request.user.pk and obj.is_active:
                    return JsonResponse({"detail": "No puedes desactivar tu propio usuario."}, status=400)
                obj.is_active = not obj.is_active
                update_fields = ["is_active"] + (["updated_at"] if hasattr(obj, "updated_at") else [])
                obj.save(update_fields=update_fields)
                detail = "Registro activado." if obj.is_active else "Registro desactivado."
                _audit_admin_action(request, AuditLog.Action.ADMIN_TOGGLED, resource, obj, {"is_active": obj.is_active})
                return JsonResponse({"detail": detail})
        if resource == "users" and obj.pk == request.user.pk:
            return JsonResponse({"detail": "No puedes eliminar tu propio usuario."}, status=400)
        try:
            deleted_metadata = {"object_id": obj.pk, "object_label": _audit_object_label(obj), "resource": resource, "resource_label": ADMIN_RESOURCE_LABELS.get(resource, resource)}
            obj.delete()
            log_action(AuditLog.Action.ADMIN_DELETED, actor=request.user, metadata=deleted_metadata, ip_address=_client_ip(request))
            return JsonResponse({"detail": "Registro eliminado."})
        except ProtectedError as error:
            return JsonResponse({"detail": _protected_dependency_message(error)}, status=400)
    payload = _json_body(request)
    if payload is None:
        return HttpResponseBadRequest("JSON inválido.")
    try:
        for field in fields:
            if field in payload:
                setattr(obj, field, payload[field])
        if resource == "users" and payload.get("password"):
            obj.set_password(payload["password"])
        if resource == "users" and "roles" in payload:
            if not payload["roles"]:
                raise ValidationError("Selecciona al menos un rol.")
            obj.rol = payload["roles"][0]
        obj.full_clean()
        _validate_administration_active_references(resource, obj)
        if resource == "users" and "campus_ids" in payload:
            _validate_user_campus_ids(obj, payload["campus_ids"])
        obj.save()
        if resource == "users" and "campus_ids" in payload:
            _set_user_campuses(obj, payload["campus_ids"])
        if resource == "users" and "roles" in payload:
            _set_user_roles(obj, payload["roles"])
        if resource == "admin-alerts" and "recipient_ids" in payload:
            _set_alert_recipients(obj, payload["recipient_ids"])
        if resource == "sections":
            _set_section_stage_responsibles(obj, payload)
        _audit_admin_action(request, AuditLog.Action.ADMIN_UPDATED, resource, obj, {"changed_fields": sorted(field for field in fields if field in payload)})
        return JsonResponse(_admin_record(resource, obj))
    except (ValidationError, IntegrityError, ValueError, TypeError) as error:
        return _validation_error(error)


# Lista secciónes visibles para el usuario autenticado.
@login_required
@require_GET
def section_list(request):
    queryset = apply_section_filters(sections_for_user(request.user), request.GET).prefetch_related(
        Prefetch(
            "stages",
            queryset=ProcessStage.objects.only("id", "section_id", "order", "status").order_by("order"),
            to_attr="progress_stages",
        )
    )
    total_count = queryset.count()
    page = None
    page_size = 10
    if request.GET.get("page"):
        try:
            page = max(1, int(request.GET["page"]))
        except (TypeError, ValueError):
            page = 1
        start = (page - 1) * page_size
        sections = queryset[start:start + page_size]
    else:
        sections = queryset[:100]
    data = []
    for section in sections:
        section_data = section_to_dict(section)
        current_stage = next(
            (stage for stage in section.progress_stages if stage.status not in {ProcessStage.Status.APPROVED, ProcessStage.Status.FINISHED}),
            section.progress_stages[-1] if section.progress_stages else None,
        )
        section_data["current_stage_order"] = current_stage.order if current_stage else 0
        data.append(section_data)
    response = {"count": total_count, "results": data}
    if page is not None:
        response.update({"page": page, "page_size": page_size, "total_pages": max(1, (total_count + page_size - 1) // page_size)})
    return JsonResponse(response)


@login_required
@require_GET
def section_filter_options(request):
    visible = sections_for_user(request.user)
    campuses = Campus.objects.filter(campus_careers__api_sections__in=visible).distinct().order_by("name")
    careers = Career.objects.filter(campus_careers__api_sections__in=visible).distinct().order_by("name")
    courses = Course.objects.filter(api_sections__in=visible).distinct().order_by("name")
    semesters = Semester.objects.filter(api_sections__in=visible).distinct().order_by("-year", "period")
    return JsonResponse({
        "campuses": [{"value": item.id, "label": item.name} for item in campuses],
        "careers": [{"value": item.id, "label": f"{item.code} · {item.name}"} for item in careers],
        "courses": [{"value": item.id, "label": f"{item.code} · {item.name}"} for item in courses],
        "semesters": [{"value": item.id, "label": str(item)} for item in semesters],
    })


def _section_detail_queryset():
    return APISection.objects.select_related(
        "campus_career__campus",
        "campus_career__career",
        "course",
        "semester",
        "career_director",
        "specialist_teacher",
        "ctp_coordinator",
        "productive_partner",
    ).prefetch_related(
        "stages__documents__uploaded_by",
        "stages__documents__reviewed_by",
        "stages__documents__approved_by",
        "stages__documents__versions__uploaded_by",
        "visits__assigned_teacher",
        "alerts__recipients",
        "ctp_reviews__submitted_by",
        "ctp_reviews__reviewed_by",
    )


def _load_section_detail(section_id):
    return get_object_or_404(_section_detail_queryset(), pk=section_id)


def _section_detail_data(request, section):
    data = section_to_dict(section, include_children=True)
    stages = list(section.stages.all())
    stages_by_id = {stage.id: stage for stage in stages}
    stage_documents = {stage.id: list(stage.documents.all()) for stage in stages}
    documents = {document.id: document for docs in stage_documents.values() for document in docs}
    document_versions = {document.id: list(document.versions.all()) for document in documents.values()}
    reviews = list(section.ctp_reviews.all())
    completed_statuses = {ProcessStage.Status.APPROVED, ProcessStage.Status.FINISHED}
    completed_orders = {stage.order for stage in stages if stage.status in completed_statuses}
    pending_review = any(review.result == CTPValidationReview.Result.PENDING for review in reviews)
    no_previous_reviews = not reviews
    stage6_rejected = any(stage.order == 6 and stage.status == ProcessStage.Status.REJECTED for stage in stages)
    visits_count = len(list(section.visits.all()))

    teachers = User.objects.filter(Q(rol=User.Role.DE) | Q(groups__name=User.Role.DE), is_active=True).distinct()
    data["teacher_options"] = [{"id": user.id, "name": user.get_full_name() or user.username} for user in teachers]
    data["validation_reviews"] = [{
        "id": review.id, "attempt": review.attempt, "result": review.result,
        "result_display": review.get_result_display(), "submitted_at": review.submitted_at.isoformat(),
        "submitted_by": review.submitted_by.get_full_name() or review.submitted_by.username if review.submitted_by else "Sistema",
        "reviewed_at": review.reviewed_at.isoformat() if review.reviewed_at else None,
        "reviewed_by": (review.reviewed_by.get_full_name() or review.reviewed_by.username) if review.reviewed_by else None,
        "feedback": review.feedback, "error_stages": review.error_stages, "details": review.details,
    } for review in reviews]

    for stage_data in data["stages"]:
        stage = stages_by_id[stage_data["id"]]
        docs = stage_documents.get(stage.id, [])
        previous_complete = all(order in completed_orders for order in range(1, stage.order))
        stage_data["is_unlocked"] = previous_complete
        if stage.order == 7:
            has_required_input = visits_count >= 3
        elif stage.order in {8, 9}:
            has_required_input = True
        else:
            has_required_input = bool(docs) and all(document_versions.get(document.id) for document in docs)
        is_closed = stage.status in completed_statuses
        stage_data["can_advance"] = stage.order in {1, 2, 3, 4, 7, 8, 9} and can_advance_stage(request.user, stage) and not is_closed
        stage_data["advance_ready"] = previous_complete and has_required_input
        stage_data["can_decide"] = stage.order == 5 and has_role(request.user, {User.Role.ADMIN, User.Role.DC}) and not is_closed
        stage_data["decision_ready"] = previous_complete and any(document.status in {Document.Status.UPLOADED, Document.Status.UNDER_REVIEW} and document_versions.get(document.id) for document in docs)
        if stage.order == 5:
            corrections_complete = all(order in completed_orders for order in range(1, 5))
            stage_five_ready = stage.status in completed_statuses or any(document.status in {Document.Status.UPLOADED, Document.Status.UNDER_REVIEW} for document in docs)
            stage_data["can_decide"] = has_role(request.user, {User.Role.ADMIN, User.Role.DC}) and not pending_review
            stage_data["can_resubmit"] = stage6_rejected and corrections_complete and stage_five_ready
            stage_data["can_submit_existing"] = no_previous_reviews and corrections_complete and stage.status in completed_statuses
            stage_data["decision_ready"] = stage_data["decision_ready"] or stage_data["can_resubmit"] or stage_data["can_submit_existing"]
            stage_data["can_reject_formulation"] = any(document.status in {Document.Status.UPLOADED, Document.Status.UNDER_REVIEW} for document in docs)
            stage_data["can_download_formulation_zip"] = has_role(request.user, {User.Role.ADMIN, User.Role.DC}) and can_access_section(request.user, section) and any(document_versions.get(document.id) for document in docs)
        stage_data["can_validate"] = stage.order == 6 and has_role(request.user, {User.Role.ADMIN, User.Role.CTP}) and not is_closed
        stage_data["validation_ready"] = previous_complete and pending_review
        stage_data["can_manage_visits"] = stage.order == 7 and can_manage_visits(request.user, section)
        stage_data["visits_unlocked"] = previous_complete
        stage_data["can_download_initial_excel"] = stage.order == 8 and has_role(request.user, {User.Role.ADMIN, User.Role.DC})
        stage_data["can_download_closure_certificate"] = stage.order == 9 and has_role(request.user, {User.Role.ADMIN, User.Role.DC})
        if stage.feedback and stage.order == 5:
            reviewed_documents = [document for document in docs if document.reviewed_by_id]
            reviewed_document = max(reviewed_documents, key=lambda document: document.reviewed_at or document.updated_at, default=None)
            if reviewed_document:
                stage_data["feedback_author"] = reviewed_document.reviewed_by.get_full_name() or reviewed_document.reviewed_by.username
                stage_data["feedback_at"] = reviewed_document.reviewed_at.isoformat() if reviewed_document.reviewed_at else None
        if stage.feedback and stage.order == 6:
            rejected_reviews = [review for review in reviews if review.result == CTPValidationReview.Result.REJECTED]
            latest_resolution = max(rejected_reviews, key=lambda review: review.attempt, default=None)
            if latest_resolution:
                stage_data["feedback_author"] = latest_resolution.reviewed_by.get_full_name() or latest_resolution.reviewed_by.username
                stage_data["feedback_at"] = latest_resolution.reviewed_at.isoformat() if latest_resolution.reviewed_at else None
        for document_data in stage_data.get("documents", []):
            document = documents[document_data["id"]]
            document_data["can_upload"] = can_upload_document(request.user, document)
            document_data["can_review"] = can_review_document(request.user, document)
    return data


def _section_detail_response(request, section_id):
    section = _load_section_detail(section_id)
    if not can_access_section(request.user, section):
        return _forbidden()
    return JsonResponse(_section_detail_data(request, section))


# Entrega el detalle completo de una sección API.
@login_required
@require_GET
def section_detail(request, section_id):
    return _section_detail_response(request, section_id)


# Reúne los documentos aprobados del ciclo en una estructura de descarga para Admin y DC.
@login_required
@require_GET
def section_cycle_download(request, section_id):
    section = get_object_or_404(
        APISection.objects.select_related("course", "campus_career__campus", "campus_career__career"),
        pk=section_id,
    )
    if not has_role(request.user, {User.Role.ADMIN, User.Role.DC}) or not can_access_section(request.user, section):
        return _forbidden()
    initial_documents = list(
        Document.objects.filter(stage__section=section, stage__order__in=[1, 2, 3, 4])
        .select_related("stage").prefetch_related("versions").order_by("stage__order", "id")
    )
    current_versions = [(document, _latest_available_version(document)) for document in initial_documents]
    current_versions = [(document, version) for document, version in current_versions if version]
    if {document.stage.order for document, version in current_versions} != {1, 2, 3, 4}:
        return _validation_error(ValidationError("Las etapas 1 a 4 deben tener sus documentos actuales cargados antes de descargar el ciclo."))

    formulation_documents = list(
        Document.objects.filter(
            stage__section=section,
            stage__order=5,
            stage__status__in=[ProcessStage.Status.APPROVED, ProcessStage.Status.FINISHED],
            status=Document.Status.APPROVED,
        ).select_related("stage").prefetch_related("versions")
    )
    formulation_versions = [
        (document, version)
        for document in formulation_documents
        for version in document.versions.order_by("version_number")
        if version.file and version.file.storage.exists(version.file.name)
    ]
    if not formulation_versions:
        return _validation_error(ValidationError("La Formulación del Proyecto debe estar aprobada por el Director de Carrera antes de descargar el ciclo."))

    try:
        consolidated = _build_initial_documents_excel(section)
    except ValidationError as error:
        return _validation_error(error)

    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        for document, version in current_versions:
            extension = version.file.name.rsplit(".", 1)[-1]
            safe_title = document.title.replace("/", "-").replace("\\", "-")
            with version.file.open("rb") as source:
                archive.writestr(f"Etapa_{document.stage.order}_{safe_title}.{extension}", source.read())
        for document, version in formulation_versions:
            original_name = version.file.name.rsplit("/", 1)[-1]
            archive_name = f"Formulación del Proyecto/{version.version_number:02d}_{original_name}"
            with version.file.open("rb") as source:
                archive.writestr(archive_name, source.read())
        archive.writestr("Excel_consolidado_etapas_1_a_3.xlsx", consolidated.getvalue())
    buffer.seek(0)
    filename = f"{section.course.code}_{section.name}_ciclo_completo.zip"
    return FileResponse(buffer, as_attachment=True, filename=filename)


def _person_name(user):
    if not user:
        return "Sin asignar"
    return user.get_full_name() or user.username


def _closure_stage_result(stage, section):
    documents = list(stage.documents.all())
    version_count = sum(document.versions.count() for document in documents)
    if stage.order in {1, 2, 3, 4}:
        titles = ", ".join(document.title for document in documents) or "Sin documento"
        return f"{titles}. {version_count} archivo(s) registrado(s)."
    if stage.order == 5:
        return f"{version_count} archivo(s) de formulación aprobados por el Director de Carrera."
    if stage.order == 6:
        reviews = list(section.ctp_reviews.all())
        approved = next((review for review in reviews if review.result == CTPValidationReview.Result.APPROVED), None)
        result = approved.get_result_display() if approved else "Sin resolución aprobatoria"
        return f"{len(reviews)} revisión(es) CTP. Resultado final: {result}."
    if stage.order == 7:
        visits = list(section.visits.all())
        visit_dates = ", ".join(visit.scheduled_date.strftime("%d-%m-%Y") for visit in visits)
        return f"{len(visits)} visita(s) registrada(s){f': {visit_dates}' if visit_dates else '.'}"
    if stage.order == 8:
        return "Excel consolidado de las etapas 1 a 3 generado y entrega final completada."
    if stage.order == 9:
        return "Cierre formal del ciclo API confirmado."
    return "Etapa completada."


@login_required
@require_GET
def section_closure_certificate(request, section_id):
    section = get_object_or_404(
        APISection.objects.select_related(
            "course", "semester", "campus_career__campus", "campus_career__career",
            "career_director", "specialist_teacher", "ctp_coordinator", "productive_partner",
        ).prefetch_related(
            "stages__documents__versions",
            "visits__assigned_teacher", "ctp_reviews__reviewed_by",
        ),
        pk=section_id,
    )
    if not has_role(request.user, {User.Role.ADMIN, User.Role.DC}) or not can_access_section(request.user, section):
        return _forbidden()
    stages = list(section.stages.all().order_by("order"))
    closure_stage = next((stage for stage in stages if stage.order == 9), None)
    if not closure_stage or closure_stage.status != ProcessStage.Status.FINISHED or section.status != APISection.Status.FINISHED:
        return _validation_error(ValidationError("El acta final estará disponible cuando se finalice la etapa 9."))

    from xml.sax.saxutils import escape as xml_escape

    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    def paragraph(value, style):
        return Paragraph(xml_escape(str(value or "")), style)

    def closure_date(value):
        return timezone.localtime(value).strftime("%d-%m-%Y %H:%M") if value else "Sin fecha"

    specialist_stage = next((stage for stage in stages if stage.order == 5), None)
    ctp_stage = next((stage for stage in stages if stage.order == 6), None)
    director_name = _person_name(section.career_director)
    specialist_name = _person_name(section.specialist_teacher)
    ctp_name = _person_name(section.ctp_coordinator)
    document_count = Document.objects.filter(stage__section=section, versions__isnull=False).distinct().count()
    file_count = DocumentVersion.objects.filter(document__stage__section=section).count()
    visit_count = section.visits.count()

    buffer = io.BytesIO()
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="ActaTitle", parent=styles["Title"], fontName="Helvetica-Bold", fontSize=17, leading=21, alignment=TA_CENTER, textColor=colors.HexColor("#C8102E"), spaceAfter=12))
    styles.add(ParagraphStyle(name="ActaSubtitle", parent=styles["Normal"], fontSize=10, leading=14, alignment=TA_CENTER, textColor=colors.HexColor("#4B5563"), spaceAfter=16))
    styles.add(ParagraphStyle(name="ActaBody", parent=styles["BodyText"], fontSize=9, leading=13, spaceAfter=7))
    styles.add(ParagraphStyle(name="ActaCell", parent=styles["BodyText"], fontSize=7.2, leading=9))
    styles.add(ParagraphStyle(name="ActaCellCenter", parent=styles["ActaCell"], alignment=TA_CENTER))
    styles.add(ParagraphStyle(name="ActaSection", parent=styles["Heading2"], fontName="Helvetica-Bold", fontSize=11, leading=14, textColor=colors.HexColor("#0E4684"), spaceBefore=7, spaceAfter=7))

    document = SimpleDocTemplate(buffer, pagesize=A4, rightMargin=14 * mm, leftMargin=14 * mm, topMargin=15 * mm, bottomMargin=15 * mm, title=f"Acta final - {section.course.code} - {section.name}", author="APIConnect")
    story = [
        Paragraph("ACTA FINAL DE CIERRE DEL CICLO API", styles["ActaTitle"]),
        Paragraph("Comprobante institucional generado automáticamente por APIConnect", styles["ActaSubtitle"]),
        Paragraph(
            f"Por medio de la presente se confirma el cierre del ciclo API de la sección <b>{xml_escape(section.name)}</b>, "
            f"correspondiente a la asignatura <b>{xml_escape(section.course.name)}</b>, carrera "
            f"<b>{xml_escape(section.campus_career.career.name)}</b>, sede <b>{xml_escape(section.campus_career.campus.name)}</b>.",
            styles["ActaBody"],
        ),
    ]

    general_data = [
        ["Código / sección", f"{section.course.code} / {section.name}", "Estado final", section.status],
        ["Asignatura", section.course.name, "Semestre", str(section.semester)],
        ["Carrera", section.campus_career.career.name, "Sede", section.campus_career.campus.name],
        ["Contraparte", section.productive_partner.name if section.productive_partner else "Sin contraparte", "Fecha de cierre", closure_date(closure_stage.completed_at)],
        ["Director de Carrera", director_name, "Docente Especialista", specialist_name],
        ["CTP responsable", ctp_name, "Registros", f"{document_count} documento(s), {file_count} archivo(s), {visit_count} visita(s)"],
    ]
    general_table = Table([[paragraph(cell, styles["ActaCell"]) for cell in row] for row in general_data], colWidths=[34 * mm, 62 * mm, 32 * mm, 54 * mm])
    general_table.setStyle(TableStyle([
        ("GRID", (0, 0), (-1, -1), 0.45, colors.HexColor("#CBD1D8")),
        ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#EEF1F5")),
        ("BACKGROUND", (2, 0), (2, -1), colors.HexColor("#EEF1F5")),
        ("FONTNAME", (0, 0), (0, -1), "Helvetica-Bold"),
        ("FONTNAME", (2, 0), (2, -1), "Helvetica-Bold"),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 5), ("RIGHTPADDING", (0, 0), (-1, -1), 5),
        ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
    ]))
    story.extend([general_table, Paragraph("Resumen de las etapas", styles["ActaSection"])])

    stage_rows = [["N°", "Etapa", "Responsable", "Estado", "Finalización", "Resultado principal"]]
    for stage in stages:
        responsible = _person_name(stage.responsible_user) if stage.responsible_user else stage.get_responsible_role_display()
        stage_rows.append([
            str(stage.order), stage.get_stage_type_display(), responsible, stage.status,
            closure_date(stage.completed_at), _closure_stage_result(stage, section),
        ])
    stage_table = Table(
        [[paragraph(cell, styles["ActaCellCenter"] if index in {0, 3, 4} else styles["ActaCell"]) for index, cell in enumerate(row)] for row in stage_rows],
        colWidths=[9 * mm, 37 * mm, 32 * mm, 20 * mm, 27 * mm, 57 * mm], repeatRows=1,
    )
    stage_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0E4684")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#CBD1D8")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F7F8FA")]),
        ("LEFTPADDING", (0, 0), (-1, -1), 3), ("RIGHTPADDING", (0, 0), (-1, -1), 3),
        ("TOPPADDING", (0, 0), (-1, -1), 4), ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]))
    story.extend([stage_table, Spacer(1, 10 * mm), Paragraph("Constancia y firmas", styles["ActaSection"]), Paragraph("Las siguientes firmas dejan constancia de la revisión y cierre del ciclo API.", styles["ActaBody"])])

    signatures = [
        [Paragraph("<br/><br/><br/>____________________________", styles["ActaCellCenter"]), Paragraph("<br/><br/><br/>____________________________", styles["ActaCellCenter"]), Paragraph("<br/><br/><br/>____________________________", styles["ActaCellCenter"])],
        [paragraph(director_name, styles["ActaCellCenter"]), paragraph(specialist_name, styles["ActaCellCenter"]), paragraph(ctp_name, styles["ActaCellCenter"])],
        [paragraph("Director de Carrera", styles["ActaCellCenter"]), paragraph("Docente Especialista", styles["ActaCellCenter"]), paragraph("CTP", styles["ActaCellCenter"])],
    ]
    signature_table = Table(signatures, colWidths=[60 * mm, 60 * mm, 60 * mm])
    signature_table.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "BOTTOM"), ("ALIGN", (0, 0), (-1, -1), "CENTER")]))
    story.extend([signature_table, Spacer(1, 8 * mm), Paragraph(f"Acta generada el {closure_date(timezone.now())}.", styles["ActaSubtitle"])])

    document.build(story)
    buffer.seek(0)
    filename = f"acta_final_{section.course.code}_{section.name}.pdf"
    return FileResponse(buffer, as_attachment=True, filename=filename, content_type="application/pdf")


# Permite al DC o rol superior definir los requisitos de una sección.
@login_required
@require_POST
def section_requirements_update(request, section_id):
    section = get_object_or_404(APISection, pk=section_id)
    if not can_define_requirements(request.user, section):
        return _forbidden()
    payload = _json_body(request)
    if payload is None:
        return HttpResponseBadRequest("JSON inválido.")
    requirements = payload.get("requirements")
    if not isinstance(requirements, str) or not requirements.strip():
        return HttpResponseBadRequest("El campo 'requirements' es obligatorio.")
    try:
        section = update_section_requirements(section, requirements, actor=request.user, ip_address=_client_ip(request))
    except ValidationError as error:
        return _validation_error(error)
    return JsonResponse(section_to_dict(section))


# Permite al DVCM o rol superior registrar una empresa propuesta.
@login_required
@require_POST
def section_partner_propose(request, section_id):
    section = get_object_or_404(APISection, pk=section_id)
    if not can_propose_partner(request.user, section):
        return _forbidden()
    payload = _json_body(request)
    if payload is None:
        return HttpResponseBadRequest("JSON inválido.")
    if not payload.get("name"):
        return HttpResponseBadRequest("El campo 'name' es obligatorio.")
    try:
        partner = propose_partner_for_section(section, payload, actor=request.user, ip_address=_client_ip(request))
    except ValidationError as error:
        return _validation_error(error)
    return JsonResponse({"partner_id": partner.id, "section": section_to_dict(section)}, status=201)


# Permite al DC o rol superior seleccionar la contraparte definitiva.
@login_required
@require_POST
def section_partner_select(request, section_id):
    section = get_object_or_404(APISection, pk=section_id)
    if not can_select_partner(request.user, section):
        return _forbidden()
    payload = _json_body(request)
    if payload is None:
        return HttpResponseBadRequest("JSON inválido.")
    partner_id = payload.get("partner_id")
    if not partner_id:
        return HttpResponseBadRequest("El campo 'partner_id' es obligatorio.")
    partner = get_object_or_404(ProductivePartner, pk=partner_id)
    try:
        section = select_partner_for_section(section, partner, actor=request.user, ip_address=_client_ip(request))
    except ValidationError as error:
        return _validation_error(error)
    return JsonResponse(section_to_dict(section))


# Lista las etapas BPM de una sección visible para el usuario.
@login_required
@require_GET
def section_stages(request, section_id):
    section = get_object_or_404(APISection, pk=section_id)
    if not can_access_section(request.user, section):
        return _forbidden()
    data = [stage_to_dict(stage, include_documents=True) for stage in section.stages.prefetch_related("documents__versions")]
    return JsonResponse({"count": len(data), "results": data})


# Lista visitas referenciales asociadas a una sección visible.
@login_required
@require_GET
def section_visits(request, section_id):
    section = get_object_or_404(APISection, pk=section_id)
    if not can_access_section(request.user, section):
        return _forbidden()
    data = [visit_to_dict(visit) for visit in section.visits.all()]
    return JsonResponse({"count": len(data), "results": data})


# Lista eventos reales del calendario desde etapas, documentos, visitas y eventos manuales.
@login_required
@require_http_methods(["GET", "POST"])
def calendar_event_list(request):
    if request.method == "POST":
        payload = _json_body(request)
        if payload is None:
            return HttpResponseBadRequest("JSON inválido.")
        section = None
        if payload.get("section_id"):
            section = get_object_or_404(APISection, pk=payload.get("section_id"))
            if not can_access_section(request.user, section):
                return _forbidden()
        try:
            event = CalendarEvent(
                title=(payload.get("title") or "").strip(),
                event_date=_parse_date(payload.get("date")),
                start_time=_parse_optional_time(payload.get("start_time")),
                end_time=_parse_optional_time(payload.get("end_time")),
                category=payload.get("category") or CalendarEvent.Category.ACTIVITY,
                section=section,
                location=(payload.get("location") or "").strip(),
                description=(payload.get("description") or "").strip(),
                created_by=request.user,
            )
            event.full_clean()
            event.save()
        except ValidationError as error:
            return _validation_error(error)
        _audit_calendar_action(request, AuditLog.Action.CALENDAR_CREATED, event)
        return JsonResponse(_custom_calendar_event_to_dict(event, request.user), status=201)

    sections = sections_for_user(request.user).prefetch_related("stages__documents", "visits")
    visible_section_ids = [section.id for section in sections]
    events = []
    for section in sections:
        for stage in section.stages.all():
            events.append(
                _calendar_event(
                    f"stage-{stage.id}",
                    "stage_due",
                    stage.get_stage_type_display(),
                    stage.due_date,
                    section,
                    stage.status,
                    {"stage_id": stage.id, "editable": can_manage_stage(request.user, stage)},
                )
            )
            for document in stage.documents.all():
                events.append(
                    _calendar_event(
                        f"document-delivery-{document.id}",
                        "document_delivery",
                        f"Entrega: {document.title}",
                        document.delivery_due_date,
                        section,
                        document.status,
                        {"document_id": document.id, "editable": can_manage_stage(request.user, stage)},
                    )
                )
                events.append(
                    _calendar_event(
                        f"document-review-{document.id}",
                        "document_review",
                        f"Revisión: {document.title}",
                        document.review_due_date,
                        section,
                        document.status,
                        {"document_id": document.id, "editable": can_manage_stage(request.user, stage)},
                    )
                )
        for visit in section.visits.all():
            events.append(
                _calendar_event(
                    f"visit-{visit.id}",
                    "visit",
                    f"Visita {visit.visit_number}",
                    visit.scheduled_date,
                    section,
                    visit.status,
                    {"visit_id": visit.id, "notes": visit.notes, "editable": can_manage_section(request.user, section)},
                )
            )
    custom_events = CalendarEvent.objects.select_related("section__course", "section__campus_career__campus", "section__campus_career__career", "created_by").filter(
        Q(section_id__in=visible_section_ids) | Q(section__isnull=True) | Q(created_by=request.user)
    )
    events.extend(_custom_calendar_event_to_dict(event, request.user) for event in custom_events)
    events = [event for event in events if event["date"]]
    events.sort(key=lambda event: (event["date"], event["metadata"].get("start_time", ""), event["title"]))
    return JsonResponse({"count": len(events), "results": events})


# Crea una visita desde el calendario para una sección visible y administrable.
@login_required
@require_POST
def calendar_visit_create(request):
    payload = _json_body(request)
    if payload is None:
        return HttpResponseBadRequest("JSON inválido.")
    section = get_object_or_404(APISection, pk=payload.get("section_id"))
    if not can_manage_visits(request.user, section):
        return _forbidden()
    try:
        visit = Visit(
            section=section,
            visit_number=int(payload.get("visit_number")),
            scheduled_date=_parse_date(payload.get("scheduled_date")),
            status=payload.get("status") or Visit.Status.SCHEDULED,
            notes=payload.get("notes", ""),
            place=payload.get("place", ""),
            assigned_teacher_id=payload.get("assigned_teacher_id") or None,
        )
        visit.full_clean()
        visit.save()
    except (TypeError, ValueError):
        return HttpResponseBadRequest("El número de visita debe ser 1, 2 o 3.")
    except IntegrityError:
        return JsonResponse({"detail": "Ya existe una visita con ese número para la sección."}, status=400)
    except ValidationError as error:
        return _validation_error(error)
    log_action(AuditLog.Action.CALENDAR_CREATED, actor=request.user, section=section, metadata={"event_type": "visit", "visit_id": visit.id, "object_label": str(visit), "event_date": visit.scheduled_date.isoformat()}, ip_address=_client_ip(request))
    return JsonResponse(visit_to_dict(visit), status=201)


# Edita o elimina un evento existente del calendario sin crear un modelo redundante.
@login_required
@require_http_methods(["POST", "PATCH", "DELETE"])
def calendar_event_update(request, event_id):
    event_type, raw_object_id = event_id.rsplit("-", 1)
    try:
        object_id = int(raw_object_id)
    except ValueError:
        return HttpResponseBadRequest("Identificador de evento inválido.")

    if request.method == "DELETE":
        if event_type == "stage":
            stage = get_object_or_404(ProcessStage.objects.select_related("section"), pk=object_id)
            if not can_manage_stage(request.user, stage):
                return _forbidden()
            stage.due_date = None
            stage.full_clean()
            stage.save(update_fields=["due_date", "updated_at"])
            log_action(AuditLog.Action.CALENDAR_DELETED, actor=request.user, section=stage.section, stage=stage, metadata={"event_type": "stage_due", "stage_id": stage.id, "object_label": stage.get_stage_type_display()}, ip_address=_client_ip(request))
            return JsonResponse({"detail": "Fecha de etapa eliminada del calendario."})

        if event_type == "document-delivery":
            document = get_object_or_404(Document.objects.select_related("stage__section"), pk=object_id)
            if not can_manage_stage(request.user, document.stage):
                return _forbidden()
            document.delivery_due_date = None
            document.full_clean()
            document.save(update_fields=["delivery_due_date", "updated_at"])
            log_action(AuditLog.Action.CALENDAR_DELETED, actor=request.user, section=document.stage.section, stage=document.stage, document=document, metadata={"event_type": "document_delivery", "document_id": document.id, "object_label": document.title}, ip_address=_client_ip(request))
            return JsonResponse({"detail": "Fecha de entrega eliminada del calendario."})

        if event_type == "document-review":
            document = get_object_or_404(Document.objects.select_related("stage__section"), pk=object_id)
            if not can_manage_stage(request.user, document.stage):
                return _forbidden()
            document.review_due_date = None
            document.full_clean()
            document.save(update_fields=["review_due_date", "updated_at"])
            log_action(AuditLog.Action.CALENDAR_DELETED, actor=request.user, section=document.stage.section, stage=document.stage, document=document, metadata={"event_type": "document_review", "document_id": document.id, "object_label": document.title}, ip_address=_client_ip(request))
            return JsonResponse({"detail": "Fecha de revisión eliminada del calendario."})

        if event_type == "visit":
            visit = get_object_or_404(Visit.objects.select_related("section"), pk=object_id)
            if not can_manage_section(request.user, visit.section):
                return _forbidden()
            visit_metadata = {"event_type": "visit", "visit_id": visit.id, "object_label": str(visit), "event_date": visit.scheduled_date.isoformat()}
            visit.delete()
            log_action(AuditLog.Action.CALENDAR_DELETED, actor=request.user, section=visit.section, metadata=visit_metadata, ip_address=_client_ip(request))
            return JsonResponse({"detail": "Visita eliminada del calendario."})

        if event_type == "custom":
            event = get_object_or_404(CalendarEvent, pk=object_id)
            if event.created_by_id != request.user.id and not has_role(request.user, ADMIN_ROLES):
                return _forbidden()
            event_metadata = {"event_type": "custom", "event_id": event.id, "event_title": event.title, "event_date": event.event_date.isoformat(), "category": event.category}
            event.delete()
            log_action(AuditLog.Action.CALENDAR_DELETED, actor=request.user, section=event.section, metadata=event_metadata, ip_address=_client_ip(request))
            return JsonResponse({"detail": "Evento eliminado del calendario."})

        return HttpResponseBadRequest("Tipo de evento no soportado.")

    payload = _json_body(request)
    if payload is None:
        return HttpResponseBadRequest("JSON inválido.")

    try:
        event_date = _parse_date(payload.get("date"))
    except ValidationError as error:
        return _validation_error(error)

    try:
        if event_type == "custom":
            event = get_object_or_404(CalendarEvent.objects.select_related("section", "created_by"), pk=object_id)
            if event.created_by_id != request.user.id and not has_role(request.user, ADMIN_ROLES):
                return _forbidden()
            section = event.section
            if "section_id" in payload:
                section = None
                if payload.get("section_id"):
                    section = get_object_or_404(APISection, pk=payload.get("section_id"))
                    if not can_access_section(request.user, section):
                        return _forbidden()
            event.title = (payload.get("title") or event.title).strip()
            event.event_date = event_date
            event.start_time = _parse_optional_time(payload.get("start_time"))
            event.end_time = _parse_optional_time(payload.get("end_time"))
            event.category = payload.get("category") or event.category
            event.section = section
            event.location = (payload.get("location") or "").strip()
            event.description = (payload.get("description") or "").strip()
            event.full_clean()
            event.save(update_fields=["title", "event_date", "start_time", "end_time", "category", "section", "location", "description", "updated_at"])
            _audit_calendar_action(request, AuditLog.Action.CALENDAR_UPDATED, event, {"changed_fields": sorted(payload.keys())})
            return JsonResponse(_custom_calendar_event_to_dict(event, request.user))

        if event_type == "stage":
            stage = get_object_or_404(ProcessStage.objects.select_related("section"), pk=object_id)
            if not can_manage_stage(request.user, stage):
                return _forbidden()
            stage.due_date = event_date
            stage.full_clean()
            stage.save(update_fields=["due_date", "updated_at"])
            log_action(AuditLog.Action.CALENDAR_UPDATED, actor=request.user, section=stage.section, stage=stage, metadata={"event_type": "stage_due", "stage_id": stage.id, "object_label": stage.get_stage_type_display(), "event_date": stage.due_date.isoformat() if stage.due_date else None}, ip_address=_client_ip(request))
            return JsonResponse(stage_to_dict(stage))

        if event_type == "document-delivery":
            document = get_object_or_404(Document.objects.select_related("stage__section"), pk=object_id)
            if not can_manage_stage(request.user, document.stage):
                return _forbidden()
            document.delivery_due_date = event_date
            document.full_clean()
            document.save(update_fields=["delivery_due_date", "updated_at"])
            log_action(AuditLog.Action.CALENDAR_UPDATED, actor=request.user, section=document.stage.section, stage=document.stage, document=document, metadata={"event_type": "document_calendar_date", "document_id": document.id, "object_label": document.title, "event_date": event_date.isoformat()}, ip_address=_client_ip(request))
            return JsonResponse(document_to_dict(document))

        if event_type == "document-review":
            document = get_object_or_404(Document.objects.select_related("stage__section"), pk=object_id)
            if not can_manage_stage(request.user, document.stage):
                return _forbidden()
            document.review_due_date = event_date
            document.full_clean()
            document.save(update_fields=["review_due_date", "updated_at"])
            log_action(AuditLog.Action.CALENDAR_UPDATED, actor=request.user, section=document.stage.section, stage=document.stage, document=document, metadata={"event_type": "document_calendar_date", "document_id": document.id, "object_label": document.title, "event_date": event_date.isoformat()}, ip_address=_client_ip(request))
            return JsonResponse(document_to_dict(document))

        if event_type == "visit":
            visit = get_object_or_404(Visit.objects.select_related("section"), pk=object_id)
            if not can_manage_section(request.user, visit.section):
                return _forbidden()
            visit.scheduled_date = event_date
            visit.status = payload.get("status") or visit.status
            visit.notes = payload.get("notes", visit.notes)
            visit.full_clean()
            visit.save(update_fields=["scheduled_date", "status", "notes", "updated_at"])
            log_action(AuditLog.Action.CALENDAR_UPDATED, actor=request.user, section=visit.section, metadata={"event_type": "visit", "visit_id": visit.id, "object_label": str(visit), "event_date": visit.scheduled_date.isoformat(), "status": visit.status}, ip_address=_client_ip(request))
            return JsonResponse(visit_to_dict(visit))
    except ValidationError as error:
        return _validation_error(error)

    return HttpResponseBadRequest("Tipo de evento no soportado.")
def _stage_partial_data(request, stage):
    data = stage_to_dict(stage, include_documents=True)
    documents = list(stage.documents.all())
    versions_by_document = {document.id: list(document.versions.all()) for document in documents}
    completed_statuses = {ProcessStage.Status.APPROVED, ProcessStage.Status.FINISHED}
    section_stages = list(stage.section.stages.all())
    section_reviews = list(stage.section.ctp_reviews.all())
    previous_complete = all(
        item.status in completed_statuses for item in section_stages if item.order < stage.order
    )
    is_closed = stage.status in completed_statuses
    if stage.order in {8, 9}:
        has_required_input = True
    elif stage.order == 7:
        has_required_input = len(list(stage.section.visits.all())) >= 3
    else:
        has_required_input = bool(documents) and all(versions_by_document.get(document.id) for document in documents)
    data["is_unlocked"] = previous_complete
    data["can_advance"] = stage.order in {1, 2, 3, 4, 7, 8, 9} and can_advance_stage(request.user, stage) and not is_closed
    data["advance_ready"] = previous_complete and has_required_input
    pending_review = any(review.result == CTPValidationReview.Result.PENDING for review in section_reviews)
    data["can_decide"] = stage.order == 5 and has_role(request.user, {User.Role.ADMIN, User.Role.DC}) and not pending_review
    if stage.order == 5:
        corrections_complete = all(item.status in completed_statuses for item in section_stages if item.order < 5)
        stage6_rejected = any(item.order == 6 and item.status == ProcessStage.Status.REJECTED for item in section_stages)
        no_previous_reviews = not section_reviews
        stage_five_ready = stage.status in completed_statuses or any(document.status in {Document.Status.UPLOADED, Document.Status.UNDER_REVIEW} for document in documents)
        data["decision_ready"] = corrections_complete and any(document.status in {Document.Status.UPLOADED, Document.Status.UNDER_REVIEW} and versions_by_document.get(document.id) for document in documents)
        data["can_resubmit"] = stage6_rejected and corrections_complete and stage_five_ready
        data["can_submit_existing"] = no_previous_reviews and corrections_complete and stage.status in completed_statuses
        data["decision_ready"] = data["decision_ready"] or data["can_resubmit"] or data["can_submit_existing"]
        data["can_reject_formulation"] = any(document.status in {Document.Status.UPLOADED, Document.Status.UNDER_REVIEW} for document in documents)
        data["can_download_formulation_zip"] = has_role(request.user, {User.Role.ADMIN, User.Role.DC}) and can_access_section(request.user, stage.section) and any(versions_by_document.get(document.id) for document in documents)
    else:
        data["decision_ready"] = False
    data["can_validate"] = stage.order == 6 and has_role(request.user, {User.Role.ADMIN, User.Role.CTP}) and not is_closed
    data["validation_ready"] = previous_complete and pending_review
    data["can_manage_visits"] = stage.order == 7 and can_manage_visits(request.user, stage.section)
    data["visits_unlocked"] = previous_complete
    data["can_download_initial_excel"] = stage.order == 8 and has_role(request.user, {User.Role.ADMIN, User.Role.DC})
    data["can_download_closure_certificate"] = stage.order == 9 and has_role(request.user, {User.Role.ADMIN, User.Role.DC})
    document_map = {document.id: document for document in documents}
    for document_data in data.get("documents", []):
        document = document_map[document_data["id"]]
        document_data["can_upload"] = can_upload_document(request.user, document)
        document_data["can_review"] = can_review_document(request.user, document)
    return data


def _stage_advance_partial_response(request, section, orders, alert=None):
    stages = ProcessStage.objects.select_related("section__campus_career__campus", "section__campus_career__career", "section__course", "section__semester", "section__career_director", "section__specialist_teacher", "section__ctp_coordinator", "section__productive_partner").prefetch_related(
        "documents__uploaded_by",
        "documents__reviewed_by",
        "documents__approved_by",
        "documents__versions__uploaded_by",
        "section__stages",
        "section__visits",
        "section__ctp_reviews",
    ).filter(section=section, order__in=orders).order_by("order")
    return JsonResponse({
        "partial": True,
        "section_id": section.id,
        "section_status": section.status,
        "section_updated_at": section.updated_at.isoformat() if section.updated_at else None,
        "stages": [_stage_partial_data(request, item) for item in stages],
        "alert": alert_to_dict(alert) if alert else None,
    })

# Avanza una etapa cuando se cumplen las reglas documentales.
@login_required
@require_POST
def stage_advance(request, stage_id):
    stage = get_object_or_404(
        ProcessStage.objects.select_related(
            "section__campus_career__campus",
            "section__campus_career__career",
            "section__course",
            "section__semester",
            "section__productive_partner",
            "section__specialist_teacher",
            "section__ctp_coordinator",
        ).prefetch_related("documents__versions", "section__stages", "section__visits"),
        pk=stage_id,
    )
    if not can_advance_stage(request.user, stage):
        return _forbidden()
    try:
        if not stage.section.career_director_id:
            raise ValidationError("Asigna un Director de Carrera a la sección antes de continuar.")
        if stage.order in {5, 6}:
            raise ValidationError("Esta etapa requiere una resolución específica.")
        previous_stages = [item for item in stage.section.stages.all() if item.order < stage.order]
        if any(item.status not in {ProcessStage.Status.APPROVED, ProcessStage.Status.FINISHED} for item in previous_stages):
            raise ValidationError("Debes completar la etapa anterior antes de avanzar.")
        documents = list(stage.documents.all())
        if stage.order == 7:
            if len(list(stage.section.visits.all())) < 3:
                raise ValidationError("Debes registrar las tres visitas obligatorias antes de avanzar.")
        elif stage.order not in {8, 9}:
            if not documents or any(not list(document.versions.all()) for document in documents):
                raise ValidationError("Debes subir el documento de esta etapa antes de avanzar.")
        if documents:
            now = timezone.now()
            Document.objects.filter(pk__in=[document.pk for document in documents]).update(
                status=Document.Status.APPROVED,
                approved_by_id=request.user.id,
                approved_at=now,
                updated_at=now,
            )
        stage = advance_stage(stage, actor=request.user, check_documents=False)
        created_alert = None
        if stage.order == 4:
            recipient = stage.section.specialist_teacher
            if recipient:
                if recipient and recipient.is_active:
                    alert = Alert.objects.create(
                        section=stage.section,
                        stage=specialist_stage,
                        alert_type=Alert.AlertType.PENDING_APPROVAL,
                        severity=Alert.Severity.YELLOW,
                        title="Formulación del Proyecto disponible",
                        message="La etapa 4 fue completada. Ya puedes subir los archivos de Formulación del Proyecto en la etapa 5.",
                        due_date=specialist_stage.due_date,
                    )
                    alert.recipients.add(recipient)
                    created_alert = alert
                    created_alert = alert
    except ValidationError as error:
        return _validation_error(error)
    stage.section.refresh_from_db(fields=["status", "updated_at"])
    affected_orders = [stage.order]
    if stage.order < 9:
        affected_orders.append(stage.order + 1)
    return _stage_advance_partial_response(request, stage.section, affected_orders, created_alert)


@login_required
@require_POST
def stage_five_decision(request, stage_id):
    stage = get_object_or_404(ProcessStage.objects.select_related("section__campus_career__campus"), pk=stage_id, order=5)
    if not has_role(request.user, {User.Role.ADMIN, User.Role.DC}) or not can_access_section(request.user, stage.section):
        return _forbidden()
    if not stage.section.career_director_id:
        return _validation_error(ValidationError("La sección debe tener un Director de Carrera asignado."))
    if stage.section.stages.filter(order__lt=5).exclude(status__in=[ProcessStage.Status.APPROVED, ProcessStage.Status.FINISHED]).exists():
        return _validation_error(ValidationError("Debes completar las primeras cuatro etapas."))
    document = stage.documents.prefetch_related("versions").first()
    if not document or not document.versions.exists():
        return _validation_error(ValidationError("El Docente Especialista debe subir la ficha antes de tomar una decisión."))
    payload = _json_body(request) or {}
    approved = payload.get("approved")
    if not isinstance(approved, bool):
        return HttpResponseBadRequest("La decisión debe indicar aprobado o rechazado.")
    feedback = str(payload.get("feedback", "")).strip()
    if approved is False and not feedback:
        return _validation_error(ValidationError("La retroalimentación es obligatoria al rechazar."))
    if stage.section.ctp_reviews.filter(result=CTPValidationReview.Result.PENDING).exists():
        return _validation_error(ValidationError("Ya existe un envío pendiente de revisión por el CTP."))
    is_resubmission = approved and stage.section.stages.filter(order=6, status=ProcessStage.Status.REJECTED).exists()
    is_existing_first_submission = approved and not stage.section.ctp_reviews.exists() and stage.status in {ProcessStage.Status.APPROVED, ProcessStage.Status.FINISHED}
    if is_resubmission:
        previous_corrections_pending = stage.section.stages.filter(order__lt=5).exclude(status__in=[ProcessStage.Status.APPROVED, ProcessStage.Status.FINISHED]).exists()
        stage_five_ready = stage.status in {ProcessStage.Status.APPROVED, ProcessStage.Status.FINISHED} or document.status in {Document.Status.UPLOADED, Document.Status.UNDER_REVIEW}
        if previous_corrections_pending or not stage_five_ready:
            return _validation_error(ValidationError("Completa todas las correcciones antes de reenviar."))
    elif not is_existing_first_submission and document.status not in {Document.Status.UPLOADED, Document.Status.UNDER_REVIEW}:
        return _validation_error(ValidationError("Debe subirse una nueva versión antes de tomar otra decisión."))
    now = timezone.now()
    document.status = Document.Status.APPROVED if approved else Document.Status.REJECTED
    document.reviewed_by = request.user; document.reviewed_at = now
    document.approved_by = request.user if approved else None
    document.approved_at = now if approved else None
    document.save()
    stage.status = ProcessStage.Status.APPROVED if approved else ProcessStage.Status.REJECTED
    stage.completed_at = now if approved else None
    stage.feedback = "" if approved else feedback
    stage.save(update_fields=["status", "completed_at", "feedback", "updated_at"])
    if approved:
        resolve_stage_alerts(stage, stage.completed_at)
        last_attempt = stage.section.ctp_reviews.order_by("-attempt").values_list("attempt", flat=True).first() or 0
        CTPValidationReview.objects.create(section=stage.section, attempt=last_attempt + 1, submitted_by=request.user)
        validation_stage = stage.section.stages.filter(order=6).first()
        if validation_stage:
            validation_stage.status = ProcessStage.Status.IN_PROGRESS
            validation_stage.completed_at = None
            validation_stage.save(update_fields=["status", "completed_at", "updated_at"])
        recipients = User.objects.filter(
            pk=stage.section.ctp_coordinator_id,
            is_active=True,
        ) if stage.section.ctp_coordinator_id else User.objects.none()
        alert = Alert.objects.create(section=stage.section, stage=validation_stage, alert_type=Alert.AlertType.PENDING_APPROVAL, severity=Alert.Severity.YELLOW, title="Formulación enviada a Validación CTP" if not is_resubmission else "Correcciones reenviadas a Validación CTP", message=f"{request.user.get_full_name() or request.user.username} aprobó la Formulación del Proyecto y envió la revisión CTP número {last_attempt + 1}.")
        alert.recipients.set(recipients)
    else:
        recipients = User.objects.filter(Q(rol=User.Role.DE) | Q(groups__name=User.Role.DE)).distinct()
        alert = Alert.objects.create(section=stage.section, stage=stage, alert_type=Alert.AlertType.PENDING_APPROVAL, severity=Alert.Severity.RED, title="Formulación del Proyecto rechazada por DC", message=f"{request.user.get_full_name() or request.user.username} rechazó la formulación del proyecto. Motivo: {feedback}", due_date=stage.due_date)
        alert.recipients.set(recipients)
    return _section_detail_response(request, stage.section_id)



@login_required
@require_GET
def stage_formulation_zip(request, stage_id):
    stage = get_object_or_404(ProcessStage.objects.select_related("section__course"), pk=stage_id, order=5)
    if not has_role(request.user, {User.Role.ADMIN, User.Role.DC}) or not can_access_section(request.user, stage.section):
        return _forbidden()
    documents = Document.objects.filter(stage=stage).select_related("stage").prefetch_related("versions")
    versions = []
    for document in documents:
        for version in document.versions.order_by("version_number"):
            if version.file and version.file.storage.exists(version.file.name):
                versions.append((document, version))
    if not versions:
        return _validation_error(ValidationError("Aún no hay archivos de formulación subidos para descargar."))
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        for document, version in versions:
            original_name = version.file.name.rsplit("/", 1)[-1]
            safe_title = document.title.replace("/", "-").replace("\\", "-")
            archive_name = f"Etapa_5_Formulacion/{safe_title}_{version.version_number:02d}_{original_name}"
            with version.file.open("rb") as source:
                archive.writestr(archive_name, source.read())
    buffer.seek(0)
    filename = f"Formulacion_{stage.section.course.code}_{stage.section.name}.zip"
    return FileResponse(buffer, as_attachment=True, filename=filename)
@login_required
@require_GET
def stage_validation_zip(request, stage_id):
    stage = get_object_or_404(ProcessStage.objects.select_related("section__course"), pk=stage_id, order=6)
    if not has_role(request.user, {User.Role.ADMIN, User.Role.CTP}) or not can_access_section(request.user, stage.section):
        return _forbidden()
    documents = Document.objects.filter(
        stage__section=stage.section,
        stage__order=5,
        stage__status=ProcessStage.Status.APPROVED,
        status=Document.Status.APPROVED,
    ).select_related("stage").prefetch_related("versions")
    if not documents.exists() or any(not item.versions.exists() for item in documents):
        return _validation_error(ValidationError("Aún no están disponibles los archivos de formulación aprobados por el Director de Carrera."))
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        for document in documents:
            for version in document.versions.order_by("version_number"):
                with version.file.open("rb") as source:
                    original_name = version.file.name.rsplit("/", 1)[-1]
                    archive_name = f"Etapa_5_Formulacion_{version.version_number:02d}_{original_name}"
                    archive.writestr(archive_name, source.read())
    buffer.seek(0)
    filename = f"{stage.section.course.code}_{stage.section.name}.zip"
    return FileResponse(buffer, as_attachment=True, filename=filename)


def _build_initial_documents_excel(section, single_sheet=False):
    documents = list(
        Document.objects.filter(
            stage__section=section,
            stage__order__in=[1, 2, 3],
            stage__status__in=[ProcessStage.Status.APPROVED, ProcessStage.Status.FINISHED],
            status=Document.Status.APPROVED,
        ).select_related("stage").prefetch_related("versions").order_by("stage__order", "id")
    )
    available_versions = [(document, _latest_available_version(document)) for document in documents]
    available_versions = [(document, version) for document, version in available_versions if version]
    available_stages = {document.stage.order for document, version in available_versions}
    if available_stages != {1, 2, 3}:
        missing_stages = sorted({1, 2, 3} - available_stages)
        raise ValidationError(f"No se encontró el archivo físico aprobado de la(s) etapa(s): {', '.join(map(str, missing_stages))}. Restaura el archivo o carga una nueva versión.")

    from openpyxl import Workbook

    workbook = Workbook()
    if single_sheet:
        consolidated_sheet = workbook.active
        consolidated_sheet.title = "Consolidado"
        consolidated_sheet.sheet_properties.pageSetUpPr.fitToPage = True
        consolidated_sheet.page_setup.fitToWidth = 1
        consolidated_sheet.page_setup.fitToHeight = 0
        next_row = 1
    else:
        workbook.remove(workbook.active)
    try:
        for document, version in available_versions:
            if single_sheet:
                next_row = _append_spreadsheet_to_single_sheet(consolidated_sheet, version, document.stage.order, next_row)
            else:
                _append_spreadsheet_to_workbook(workbook, version, document.stage.order)
        buffer = io.BytesIO()
        workbook.save(buffer)
    finally:
        workbook.close()
    buffer.seek(0)
    return buffer


@login_required
@require_GET
def stage_initial_documents_excel(request, stage_id):
    stage = get_object_or_404(ProcessStage.objects.select_related("section__course"), pk=stage_id, order=8)
    if not can_access_section(request.user, stage.section) or not has_role(request.user, {User.Role.ADMIN, User.Role.DC}):
        return _forbidden()
    single_sheet = request.GET.get("layout") == "single"
    try:
        buffer = _build_initial_documents_excel(stage.section, single_sheet=single_sheet)
    except ValidationError as error:
        return _validation_error(error)
    suffix = "etapas_1_a_3_una_hoja" if single_sheet else "etapas_1_a_3"
    filename = f"{stage.section.course.code}_{stage.section.name}_{suffix}.xlsx"
    return FileResponse(buffer, as_attachment=True, filename=filename)




@login_required
@require_POST
def stage_six_decision(request, stage_id):
    stage = get_object_or_404(ProcessStage.objects.select_related("section__campus_career__campus"), pk=stage_id, order=6)
    if not has_role(request.user, {User.Role.ADMIN, User.Role.CTP}) or not can_access_section(request.user, stage.section):
        return _forbidden()
    if not stage.section.career_director_id:
        return _validation_error(ValidationError("La sección debe tener un Director de Carrera asignado para recibir la resolución."))
    review = stage.section.ctp_reviews.filter(result=CTPValidationReview.Result.PENDING).order_by("-attempt").first()
    if not review:
        return _validation_error(ValidationError("No existe un envío pendiente de revisión CTP."))
    if stage.section.stages.filter(order__lt=6).exclude(status__in=[ProcessStage.Status.APPROVED, ProcessStage.Status.FINISHED]).exists():
        return _validation_error(ValidationError("Las primeras cinco etapas deben estar aprobadas."))
    payload = _json_body(request) or {}
    approved = payload.get("approved")
    feedback = str(payload.get("feedback", "")).strip()
    if not isinstance(approved, bool):
        return HttpResponseBadRequest("La resolución enviada no es válida.")
    if approved is False and not feedback:
        return _validation_error(ValidationError("Indica la retroalimentación del proyecto de formulación."))
    project_stage = stage.section.stages.filter(order=5).first()
    if approved is False and not project_stage:
        return _validation_error(ValidationError("No existe la etapa de formulación del proyecto para retroalimentar."))
    error_stages = [] if approved else [5]
    details = {} if approved else {"5": feedback}
    now = timezone.now()
    stage.status = ProcessStage.Status.APPROVED if approved else ProcessStage.Status.REJECTED
    stage.completed_at = now if approved else None
    stage.feedback = ""
    stage.review_details = {"project_feedback": feedback, "review_date": now.date().isoformat()}
    stage.save(update_fields=["status", "completed_at", "feedback", "review_details", "updated_at"])
    if approved:
        resolve_stage_alerts(stage, stage.completed_at)
    review.result = CTPValidationReview.Result.APPROVED if approved else CTPValidationReview.Result.REJECTED
    review.reviewed_by = request.user
    review.reviewed_at = now
    review.feedback = feedback
    review.error_stages = error_stages
    review.details = details
    review.save(update_fields=["result", "reviewed_by", "reviewed_at", "feedback", "error_stages", "details", "updated_at"])
    if approved:
        recipients = User.objects.filter(pk=stage.section.career_director_id)
        title, message, severity = "Validación CTP aprobada", f"{request.user.get_full_name() or request.user.username} aprobó la formulación del proyecto. La sección puede avanzar al Registro de Visitas.", Alert.Severity.GREEN
    else:
        project_stage.status = ProcessStage.Status.REJECTED
        project_stage.completed_at = None
        project_stage.feedback = feedback
        project_stage.save(update_fields=["status", "completed_at", "feedback", "updated_at"])
        project_stage.documents.update(status=Document.Status.REJECTED, approved_by=None, approved_at=None)
        recipients = User.objects.filter(pk=stage.section.career_director_id)
        title, message, severity = "Validación CTP rechazada", f"{request.user.get_full_name() or request.user.username} rechazó la formulación del proyecto. Motivo: {feedback}", Alert.Severity.RED
    alert = Alert.objects.create(section=stage.section, stage=stage, alert_type=Alert.AlertType.PENDING_APPROVAL, severity=severity, title=title, message=message)
    alert.recipients.set(recipients)
    if not approved:
        specialist = project_stage
        specialist_recipients = User.objects.filter(pk=stage.section.specialist_teacher_id) if stage.section.specialist_teacher_id else User.objects.none()
        specialist_alert = Alert.objects.create(section=stage.section, stage=specialist, alert_type=Alert.AlertType.PENDING_APPROVAL, severity=Alert.Severity.RED, title="Corrección solicitada por CTP en Formulación del Proyecto", message=f"{request.user.get_full_name() or request.user.username} solicitó correcciones sobre la formulación del proyecto. Motivo: {feedback}")
        specialist_alert.recipients.set(specialist_recipients)
    return _section_detail_response(request, stage.section_id)


# Lista documentos visibles para el usuario autenticado.
@login_required
@require_GET
def document_list(request):
    queryset = documents_for_user(request.user).select_related(
        "stage__section__course",
        "stage__section__semester",
        "stage__section__campus_career__campus",
        "stage__section__campus_career__career",
    )
    status = request.GET.get("status")
    stage_id = request.GET.get("stage")
    if status:
        queryset = queryset.filter(status=status)
    if stage_id:
        queryset = queryset.filter(stage_id=stage_id)
    data = []
    for document in queryset[:100]:
        item = document_to_dict(document, include_versions=True)
        item["can_upload"] = can_upload_document(request.user, document)
        item["can_review"] = can_review_document(request.user, document)
        data.append(item)
    return JsonResponse({"count": queryset.count(), "results": data})


# Permite al CTP o rol superior aprobar o rechazar documentos.
@login_required
@require_POST
def document_review(request, document_id):
    document = get_object_or_404(Document.objects.select_related("stage__section"), pk=document_id)
    if not can_review_document(request.user, document):
        return _forbidden()

    payload = _json_body(request)
    if payload is None:
        return HttpResponseBadRequest("JSON inválido.")

    approved = payload.get("approved")
    if not isinstance(approved, bool):
        return HttpResponseBadRequest("El campo 'approved' debe ser booleano.")

    try:
        document = review_document(document, reviewer=request.user, approved=approved)
    except ValidationError as error:
        return _validation_error(error)
    return JsonResponse(document_to_dict(document, include_versions=True))


# Permite subir una version nueva de un documento si el rol puede gestionar la etapa.
@login_required
@require_POST
def document_upload(request, document_id):
    document = get_object_or_404(Document.objects.select_related("stage__section"), pk=document_id)
    if not can_upload_document(request.user, document):
        return _forbidden()
    uploaded_files = request.FILES.getlist("files") or request.FILES.getlist("file")
    if not uploaded_files:
        return HttpResponseBadRequest("El archivo es obligatorio.")
    if document.stage.order != 5 and len(uploaded_files) > 1:
        return _validation_error(ValidationError("Solo Formulación del Proyecto permite subir varios archivos."))
    if document.stage.order == 5:
        current_count = document.versions.count()
        if current_count + len(uploaded_files) > 15:
            available_slots = max(0, 15 - current_count)
            return _validation_error(ValidationError(f"La Formulación del Proyecto permite un máximo de 15 archivos. Ya hay {current_count} archivo(s) subido(s); puedes agregar {available_slots} más."))
    try:
        uploaded_versions = [
            add_document_version(document, uploaded_file, uploaded_by=request.user, notes=request.POST.get("notes", ""))
            for uploaded_file in uploaded_files
        ]
    except ValidationError as error:
        return _validation_error(error)
    document.refresh_from_db()
    if document.stage.order == 5 and document.stage.section.career_director_id:
        uploader_name = request.user.get_full_name() or request.user.username
        previous_versions = document.versions.count() - len(uploaded_versions)
        is_correction = previous_versions > 0
        file_count = len(uploaded_versions)
        upload_description = "subió nuevos archivos corregidos de" if is_correction and file_count > 1 else "subió una nueva versión corregida de"
        if not is_correction:
            upload_description = "subió archivos de" if file_count > 1 else "subió la"
        file_detail = f" ({file_count} archivos)" if file_count > 1 else ""
        alert = Alert.objects.create(
            section=document.stage.section,
            stage=document.stage,
            alert_type=Alert.AlertType.PENDING_APPROVAL,
            severity=Alert.Severity.YELLOW,
            title="Corrección de formulación disponible para revisión" if is_correction else "Formulación del proyecto disponible para revisión",
            message=f"{uploader_name} {upload_description} formulación del proyecto{file_detail}. Revisa la etapa 5 para aprobarla o solicitar correcciones.",
            due_date=document.stage.due_date,
        )
        alert.recipients.add(document.stage.section.career_director)
    return JsonResponse(document_to_dict(document, include_versions=True), status=201)


# Permite retirar un archivo concreto de la formulación mientras la etapa siga editable.
@login_required
@require_http_methods(["DELETE"])
def document_version_delete(request, version_id):
    version = get_object_or_404(
        DocumentVersion.objects.select_related("document__stage__section__campus_career__campus"),
        pk=version_id,
    )
    document = version.document
    if document.stage.order != 5:
        return _validation_error(ValidationError("Solo se pueden eliminar archivos de Formulación del Proyecto."))
    if not can_upload_document(request.user, document):
        return _forbidden()
    document = delete_document_version(version, deleted_by=request.user, ip_address=_client_ip(request))
    return JsonResponse({"deleted": True, "document": document_to_dict(document, include_versions=True)})


# Descarga una version de documento solo si el usuario tiene acceso a la sección.
@login_required
@require_GET
def document_download(request, version_id):
    version = get_object_or_404(
        DocumentVersion.objects.select_related("document__stage__section__campus_career__campus"),
        pk=version_id,
    )
    if not can_download_document(request.user, version.document):
        return _forbidden()
    if version.document.stage.order == 5 and not has_role(request.user, {User.Role.ADMIN, User.Role.CTP}):
        return _forbidden()
    log_action(
        AuditLog.Action.DOCUMENT_DOWNLOADED,
        actor=request.user,
        section=version.document.stage.section,
        stage=version.document.stage,
        document=version.document,
        metadata={"version": version.version_number},
        ip_address=_client_ip(request),
    )
    return FileResponse(version.file.open("rb"), as_attachment=True, filename=version.file.name.rsplit("/", 1)[-1])


# Devuelve una vista previa tabular de la version Excel subida.
@login_required
@require_GET
def document_version_preview(request, version_id):
    version = get_object_or_404(
        DocumentVersion.objects.select_related("document__stage__section__campus_career__campus"),
        pk=version_id,
    )
    if not can_download_document(request.user, version.document):
        return _forbidden()
    try:
        preview = _excel_preview(version)
    except ValidationError as error:
        return _validation_error(error)
    except Exception:
        return JsonResponse(
            {"detail": "No se pudo leer la vista previa. Descarga el archivo para revisarlo completo."},
            status=400,
        )
    return JsonResponse(
        {
            "version_id": version.id,
            "file_name": version.file.name.rsplit("/", 1)[-1],
            "sheet": preview["sheet"],
            "rows": preview["rows"],
        }
    )


# Lista alertas visibles para el usuario autenticado.
@login_required
@require_GET
def alert_list(request):
    queryset = alerts_for_user(request.user)
    if request.GET.get("active") in {"1", "true", "True"}:
        queryset = queryset.filter(is_active=True)
    severity = request.GET.get("severity")
    if severity:
        queryset = queryset.filter(severity=severity)
    data = [alert_to_dict(alert) for alert in queryset[:100]]
    return JsonResponse({"count": queryset.count(), "results": data})


# Genera alertas por vencimientos; reservado para roles administrativos.
@login_required
@require_POST
def alert_generate(request):
    if not has_role(request.user, ADMIN_ROLES):
        return _forbidden()
    alerts = generate_due_alerts()
    return JsonResponse({"created": len(alerts), "results": [alert_to_dict(alert) for alert in alerts]})










































