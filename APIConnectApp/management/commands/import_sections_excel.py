from pathlib import Path

from django.contrib.auth.models import Group
from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from openpyxl import load_workbook

from APIConnectApp.models import (
    APISection,
    Campus,
    CampusCareer,
    Career,
    Course,
    Semester,
    User,
    UserCampus,
)


EXPECTED_HEADERS = (
    "Programa de Estudio",
    "Cód Asignatura",
    "Asignatura",
    "Sección",
    "Rut Prof.",
    "Profesor",
    "Director de Carrera",
    "Correo DC",
    "CTP",
    "Correo CTP",
    "Docente Especialista",
    "Correo DE",
)


class Command(BaseCommand):
    help = "Importa secciones y responsables desde la planilla institucional."

    def add_arguments(self, parser):
        parser.add_argument("file", type=Path)
        parser.add_argument("--campus", required=True)
        parser.add_argument("--year", required=True, type=int)
        parser.add_argument("--period", required=True)

    @transaction.atomic
    def handle(self, *args, **options):
        path = options["file"].expanduser().resolve()
        if not path.is_file():
            raise CommandError(f"No existe el archivo: {path}")

        campus = self._one(Campus, "sede", name=options["campus"], is_active=True)
        semester = self._one(
            Semester,
            "semestre",
            year=options["year"],
            period=options["period"],
            is_active=True,
        )
        workbook = load_workbook(path, read_only=True, data_only=True)
        sheet = workbook.active
        headers = tuple(cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1)))
        if headers != EXPECTED_HEADERS:
            raise CommandError("Las columnas del Excel no corresponden al formato esperado.")

        created_sections = updated_sections = created_users = 0
        for row_number, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
            if not any(values):
                continue
            row = dict(zip(EXPECTED_HEADERS, values))
            required = (
                "Programa de Estudio", "Cód Asignatura", "Asignatura", "Sección",
                "Director de Carrera", "Correo DC", "CTP", "Correo CTP",
                "Docente Especialista", "Correo DE",
            )
            missing = [field for field in required if not str(row.get(field) or "").strip()]
            if missing:
                raise CommandError(f"Fila {row_number}: faltan {', '.join(missing)}.")

            career_code, career_name = self._split_program(row["Programa de Estudio"], row_number)
            career, _ = Career.objects.update_or_create(
                code=career_code,
                defaults={"name": career_name, "is_active": True},
            )
            campus_career, _ = CampusCareer.objects.update_or_create(
                campus=campus,
                career=career,
                defaults={"is_active": True},
            )
            course, _ = Course.objects.update_or_create(
                code=str(row["Cód Asignatura"]).strip(),
                defaults={"name": str(row["Asignatura"]).strip(), "is_active": True},
            )

            director, was_created = self._user(
                row["Director de Carrera"], row["Correo DC"], User.Role.DC, campus
            )
            created_users += int(was_created)
            ctp, was_created = self._user(row["CTP"], row["Correo CTP"], User.Role.CTP, campus)
            created_users += int(was_created)
            specialist, was_created = self._user(
                row["Docente Especialista"], row["Correo DE"], User.Role.DE, campus
            )
            created_users += int(was_created)

            _, was_created = APISection.objects.update_or_create(
                campus_career=campus_career,
                course=course,
                semester=semester,
                name=str(row["Sección"]).strip(),
                defaults={
                    "career_director": director,
                    "ctp_coordinator": ctp,
                    "specialist_teacher": specialist,
                },
            )
            created_sections += int(was_created)
            updated_sections += int(not was_created)

        self.stdout.write(self.style.SUCCESS(
            f"Importación completada: {created_sections} secciones creadas, "
            f"{updated_sections} actualizadas y {created_users} usuarios creados."
        ))

    @staticmethod
    def _one(model, label, **filters):
        try:
            return model.objects.get(**filters)
        except model.DoesNotExist as error:
            raise CommandError(f"No existe {label} con los datos {filters}.") from error
        except model.MultipleObjectsReturned as error:
            raise CommandError(f"Hay más de un {label} con los datos {filters}.") from error

    @staticmethod
    def _split_program(value, row_number):
        parts = [part.strip() for part in str(value).split(" - ", 1)]
        if len(parts) != 2 or not all(parts):
            raise CommandError(f"Fila {row_number}: programa de estudio inválido.")
        return parts

    @staticmethod
    def _user(name, email, role, campus):
        name = str(name).strip()
        email = str(email).strip().lower()
        user = User.objects.filter(email__iexact=email).first()
        created = user is None
        if created:
            username = name
            suffix = 2
            while User.objects.filter(username=username).exists():
                username = f"{name} {suffix}"
                suffix += 1
            user = User(username=username, email=email, rol=role, is_active=True)
            user.set_unusable_password()
            user.save()
        else:
            changed = []
            if not user.is_active:
                user.is_active = True
                changed.append("is_active")
            if user.rol != role:
                user.rol = role
                changed.append("rol")
            if changed:
                user.save(update_fields=changed)

        group, _ = Group.objects.get_or_create(name=role)
        user.groups.add(group)
        UserCampus.objects.update_or_create(
            user=user, campus=campus, defaults={"is_active": True}
        )
        return user, created
