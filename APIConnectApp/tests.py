from datetime import timedelta
import io

from django.core.exceptions import ValidationError
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from django.urls import reverse
from django.utils import timezone
from openpyxl import Workbook, load_workbook

from .models import APISection, Campus, CampusCareer, Career, Course, Document, ProcessStage, ProductivePartner, Semester, User, UserCampus
from .services import add_document_version, generate_due_alerts, review_document


# Pruebas de flujo BPM, seguridad por sede y acciones por rol.
class BPMBackendTests(TestCase):
    # Prepara una seccion base con sede, carrera, curso y semestre.
    def setUp(self):
        self.user = User.objects.create_user(username="coordinador", email="coordinador@example.com", password="secret", rol=User.Role.C_API)
        self.campus = Campus.objects.create(name="Sede Central")
        UserCampus.objects.create(user=self.user, campus=self.campus)
        self.career = Career.objects.create(code="INF", name="Informatica")
        self.campus_career = CampusCareer.objects.create(campus=self.campus, career=self.career)
        self.course = Course.objects.create(code="API101", name="Proyecto API")
        self.semester = Semester.objects.create(year=2026, period=Semester.Period.OTONO)
        self.section = APISection.objects.create(
            campus_career=self.campus_career,
            course=self.course,
            semester=self.semester,
            name="001",
            career_director=self.user,
        )

    # Verifica que una seccion cree sus etapas y documentos iniciales.
    def test_section_creation_builds_bpm_stages_and_documents(self):
        self.assertEqual(self.section.stages.count(), 9)
        first_stage = self.section.stages.get(order=1)
        self.assertEqual(first_stage.stage_type, ProcessStage.StageType.REQUIREMENTS)
        self.assertEqual(first_stage.documents.count(), 1)

    # Verifica alertas cuando una etapa esta cercana al vencimiento.
    def test_generate_due_alerts_creates_warning_alert(self):
        stage = self.section.stages.get(order=1)
        stage.due_date = timezone.localdate() + timedelta(days=1)
        stage.save()

        alerts = generate_due_alerts()

        self.assertEqual(len(alerts), 1)
        self.assertEqual(alerts[0].section, self.section)

    # Verifica que el listado requiera sesion y respete acceso del usuario.
    def test_section_list_requires_login_and_returns_accessible_sections(self):
        self.client.login(username="coordinador", password="secret")

        response = self.client.get(reverse("api_connect:section-list"))

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()["count"], 1)

    # Verifica que UserCampus limite la visibilidad por sede.
    def test_user_only_sees_assigned_campus_sections(self):
        other_campus = Campus.objects.create(name="Concepcion")
        other_campus_career = CampusCareer.objects.create(campus=other_campus, career=self.career)
        APISection.objects.create(
            campus_career=other_campus_career,
            course=self.course,
            semester=self.semester,
            name="002",
            career_director=self.user,
        )
        limited_user = User.objects.create_user(username="ctp-la", email="ctp-la@example.com", password="secret", rol=User.Role.CTP)
        UserCampus.objects.create(user=limited_user, campus=self.campus)
        self.client.login(username="ctp-la", password="secret")

        response = self.client.get(reverse("api_connect:section-list"))

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()["count"], 1)
        self.assertEqual(response.json()["results"][0]["id"], self.section.id)

    # Verifica que Coordinador API no pueda modificar requisitos.
    def test_coordinator_api_cannot_update_requirements(self):
        self.client.login(username="coordinador", password="secret")

        response = self.client.post(
            reverse("api_connect:section-requirements-update", args=[self.section.id]),
            data='{"requirements": "Nuevos requisitos"}',
            content_type="application/json",
        )

        self.assertEqual(response.status_code, 403)

    # Verifica que DVCM pueda proponer empresa, pero no seleccionarla.
    def test_dvcm_can_propose_but_not_select_partner(self):
        dvcm = User.objects.create_user(username="dvcm", email="dvcm@example.com", password="secret", rol=User.Role.D_VCM)
        UserCampus.objects.create(user=dvcm, campus=self.campus)
        self.client.login(username="dvcm", password="secret")

        response = self.client.post(
            reverse("api_connect:section-partner-propose", args=[self.section.id]),
            data='{"name": "Empresa API"}',
            content_type="application/json",
        )

        self.section.refresh_from_db()
        self.assertEqual(response.status_code, 201)
        self.assertIsNone(self.section.productive_partner)

        partner_id = response.json()["partner_id"]
        select_response = self.client.post(
            reverse("api_connect:section-partner-select", args=[self.section.id]),
            data=f'{{"partner_id": {partner_id}}}',
            content_type="application/json",
        )

        self.assertEqual(select_response.status_code, 403)

    # Verifica que DC pueda seleccionar la contraparte final.
    def test_dc_can_select_partner(self):
        dc = User.objects.create_user(username="dc", email="dc@example.com", password="secret", rol=User.Role.DC)
        UserCampus.objects.create(user=dc, campus=self.campus)
        self.section.career_director = dc
        self.section.save()
        partner = ProductivePartner.objects.create(name="Empresa seleccionada")
        self.client.login(username="dc", password="secret")

        response = self.client.post(
            reverse("api_connect:section-partner-select", args=[self.section.id]),
            data=f'{{"partner_id": {partner.id}}}',
            content_type="application/json",
        )

        self.assertEqual(response.status_code, 200)
        self.section.refresh_from_db()
        self.assertEqual(self.section.productive_partner, partner)

    # Verifica rechazo de extensiones de archivo no permitidas.
    def test_document_version_rejects_invalid_file_type(self):
        document = self.section.stages.get(order=1).documents.first()
        invalid_file = SimpleUploadedFile("payload.exe", b"bad")

        with self.assertRaises(ValidationError):
            add_document_version(document, invalid_file, uploaded_by=self.user)

    # Verifica que CTP pueda aprobar documentos con archivo versionado.
    def test_ctp_can_approve_uploaded_document(self):
        ctp = User.objects.create_user(username="ctp", email="ctp@example.com", password="secret", rol=User.Role.CTP)
        UserCampus.objects.create(user=ctp, campus=self.campus)
        document = self.section.stages.get(order=1).documents.first()
        uploaded_file = SimpleUploadedFile("ficha.xlsx", b"content")
        add_document_version(document, uploaded_file, uploaded_by=self.user)

        reviewed = review_document(document, reviewer=ctp, approved=True)

        self.assertEqual(reviewed.status, Document.Status.APPROVED)

    # El avance devuelve solo las etapas afectadas para mantener la respuesta rapida.
    def test_stage_advance_returns_partial_update(self):
        admin = User.objects.create_superuser(
            username="admin-avance",
            email="admin@example.com",
            password="secret",
            rol=User.Role.ADMIN,
        )
        document = self.section.stages.get(order=1).documents.first()
        add_document_version(
            document,
            SimpleUploadedFile(
                "ficha.xlsx",
                b"archivo-de-prueba",
                content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            ),
            uploaded_by=admin,
        )
        self.client.force_login(admin)

        response = self.client.post(reverse("api_connect:stage-advance", args=[document.stage_id]))

        self.assertEqual(response.status_code, 200)
        payload = response.json()
        self.assertTrue(payload["partial"])
        self.assertEqual([stage["order"] for stage in payload["stages"]], [1, 2])
        self.assertEqual(payload["stages"][0]["status"], ProcessStage.Status.FINISHED)

    def test_stage_responsibles_are_derived_from_section(self):
        specialist = User.objects.create_user(
            username="especialista", email="especialista@example.com", password="secret", rol=User.Role.DE
        )
        ctp = User.objects.create_user(
            username="ctp-asignado", email="ctp-asignado@example.com", password="secret", rol=User.Role.CTP
        )
        self.section.specialist_teacher = specialist
        self.section.ctp_coordinator = ctp
        self.section.save(update_fields=["specialist_teacher", "ctp_coordinator", "updated_at"])

        self.assertEqual(self.section.stages.get(order=1).responsible_user_id, self.user.id)
        self.assertEqual(self.section.stages.get(order=5).responsible_user_id, specialist.id)
        self.assertEqual(self.section.stages.get(order=6).responsible_user_id, ctp.id)

    def test_sections_bulk_import_and_export_use_institutional_format(self):
        admin = User.objects.create_superuser(
            username="admin-import", email="admin-import@example.com", password="secret", rol=User.Role.ADMIN
        )
        workbook = Workbook()
        sheet = workbook.active
        headers = [
            "Programa de Estudio", "Cód Asignatura", "Asignatura", "Sección",
            "Rut Prof.", "Profesor", "Director de Carrera", "Correo DC",
            "CTP", "Correo CTP", "Docente Especialista", "Correo DE",
        ]
        sheet.append(headers)
        sheet.append([
            "TST - Carrera de Prueba", "API900", "Proyecto de Prueba", "SEC-900",
            "1-9", "Profesor Ignorado", "Directora Prueba", "dc-import@example.com",
            "CTP Prueba", "ctp-import@example.com", "Especialista Prueba", "de-import@example.com",
        ])
        content = io.BytesIO()
        workbook.save(content)
        content.seek(0)
        upload = SimpleUploadedFile(
            "secciones.xlsx", content.read(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        self.client.force_login(admin)

        response = self.client.post(
            reverse("api_connect:section-bulk-import"),
            {"file": upload, "campus_id": self.campus.id, "semester_id": self.semester.id},
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()["created_sections"], 1)
        imported = APISection.objects.get(name="SEC-900")
        self.assertEqual(imported.specialist_teacher.email, "de-import@example.com")
        self.assertFalse(User.objects.filter(username="Profesor Ignorado").exists())

        export_response = self.client.get(reverse("api_connect:section-bulk-export"))
        self.assertEqual(export_response.status_code, 200)
        exported = load_workbook(io.BytesIO(export_response.content), read_only=True).active
        self.assertEqual(tuple(cell.value for cell in exported[1]), tuple(headers))
        exported_rows = list(exported.iter_rows(min_row=2, values_only=True))
        self.assertTrue(any(row[3] == "SEC-900" and row[10] == "Especialista Prueba" for row in exported_rows))
